#!/usr/bin/env python3
"""Render the two new slide bodies into a clean, throwaway presentation.

Slide 1 of the scratch deck = "Background"
Slide 2 of the scratch deck = "Methods overview"

Nothing here touches the user's file. build_deck_surgical.py transplants the
resulting shape trees into the original package.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

import update_deck_v2 as U   # reuse all the layout helpers

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / ".build" / "scratch_content.pptx"


def main() -> None:
    U.BUILD.mkdir(exist_ok=True)
    U.extract_media()
    U.build_panels()

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    # --- slide 1: Background ------------------------------------------------
    U.build_background_slide(prs)

    # --- slide 2: Methods overview -----------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    U.add_header(slide, "Experimental programme and evidence hierarchy",
                 "Methods overview", 7)
    U.rebuild_methods_body(slide)

    # the hero image for slide 1 of the real deck
    U.cover(U.HERO, "hero_ai", 3.63, 4.54, enhance=False)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
