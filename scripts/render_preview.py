#!/usr/bin/env python3
"""Approximate raster preview of PPTX slides, for layout QA without LibreOffice.

Renders shape rectangles, fills, pictures, tables and wrapped text at roughly the
right size so that overflow, collisions and empty regions are visible.
"""
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SCALE = 110  # px per inch
ROOT = Path(__file__).resolve().parents[1]


def font(size_pt, bold=False):
    px = max(7, int(size_pt * SCALE / 72))
    for name in (("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"),):
        for base in ("/usr/share/fonts/truetype/dejavu/",):
            try:
                return ImageFont.truetype(base + name, px)
            except OSError:
                pass
    return ImageFont.load_default()


def hexof(color_fmt, default=None):
    try:
        return "#%02X%02X%02X" % tuple(color_fmt.rgb)
    except Exception:
        return default


def wrap(draw, text, fnt, max_px):
    out = []
    for raw in text.split("\n"):
        line = ""
        for word in raw.split(" "):
            trial = (line + " " + word).strip()
            if draw.textlength(trial, font=fnt) <= max_px or not line:
                line = trial
            else:
                out.append(line)
                line = word
        out.append(line)
    return out


def draw_text_frame(draw, tf, x, y, w, h, overflow):
    cy = y + 2
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            cy += 6
            continue
        text = "".join(r.text for r in runs)
        r0 = runs[0]
        size = r0.font.size.pt if r0.font.size else 12
        bold = bool(r0.font.bold)
        col = hexof(r0.font.color, "#202B33") or "#202B33"
        fnt = font(size, bold)
        for line in wrap(draw, text, fnt, max(10, w - 6)):
            draw.text((x + 3, cy), line, font=fnt, fill=col)
            cy += int(size * SCALE / 72 * 1.18)
    if cy > y + h + 3:
        overflow.append((x, y, w, h, cy - (y + h)))
    return cy


def render(pptx_path, out_dir, only=None):
    prs = Presentation(str(pptx_path))
    W = int(prs.slide_width / 914400 * SCALE)
    H = int(prs.slide_height / 914400 * SCALE)
    out_dir.mkdir(parents=True, exist_ok=True)
    report = []
    for idx, slide in enumerate(prs.slides, start=1):
        if only and idx not in only:
            continue
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)
        overflow = []
        for sh in slide.shapes:
            if sh.left is None:
                continue
            x = int(sh.left / 914400 * SCALE)
            y = int(sh.top / 914400 * SCALE)
            w = int((sh.width or 0) / 914400 * SCALE)
            h = int((sh.height or 0) / 914400 * SCALE)
            st = str(sh.shape_type)
            if "PICTURE" in st:
                try:
                    im = Image.open(__import__("io").BytesIO(sh.image.blob)).convert("RGB")
                    img.paste(im.resize((max(1, w), max(1, h))), (x, y))
                except Exception:
                    d.rectangle([x, y, x + w, y + h], fill="#DDDDDD", outline="#888888")
                continue
            if "CHART" in st:
                d.rectangle([x, y, x + w, y + h], fill="#F2F6F9", outline="#8AA6B8")
                d.text((x + 6, y + 6), "[chart]", font=font(10), fill="#4A6B80")
                continue
            if sh.has_table:
                t = sh.table
                col_px = [int(c.width / 914400 * SCALE) for c in t.columns]
                row_px = [int(r.height / 914400 * SCALE) for r in t.rows]
                total_h = sum(row_px)
                cy = y
                for ri, row in enumerate(t.rows):
                    cx = x
                    for ci in range(len(col_px)):
                        cell = t.cell(ri, ci)
                        fill = hexof(cell.fill.fore_color, "#FFFFFF") if cell.fill.type is not None else "#FFFFFF"
                        d.rectangle([cx, cy, cx + col_px[ci], cy + row_px[ri]],
                                    fill=fill, outline="#C0CCD5")
                        txt = cell.text
                        runs = cell.text_frame.paragraphs[0].runs
                        size = runs[0].font.size.pt if runs and runs[0].font.size else 9
                        bold = bool(runs and runs[0].font.bold)
                        col = hexof(runs[0].font.color, "#202B33") if runs else "#202B33"
                        fnt = font(size, bold)
                        ty = cy + 3
                        lines = wrap(d, txt, fnt, col_px[ci] - 8)
                        for ln in lines:
                            d.text((cx + 4, ty), ln, font=fnt, fill=col or "#202B33")
                            ty += int(size * SCALE / 72 * 1.16)
                        if ty > cy + row_px[ri] + 2:
                            overflow.append(("table", ri, ci, ty - (cy + row_px[ri])))
                        cx += col_px[ci]
                    cy += row_px[ri]
                if cy > y + total_h + 2:
                    pass
                continue
            # autoshape / textbox
            fill = None
            try:
                if sh.fill.type is not None and sh.fill.type == 1:
                    fill = hexof(sh.fill.fore_color)
            except Exception:
                pass
            line = None
            try:
                line = hexof(sh.line.color)
            except Exception:
                pass
            if fill or line:
                d.rectangle([x, y, x + max(1, w), y + max(1, h)], fill=fill, outline=line)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text_frame(d, sh.text_frame, x, y, w, h, overflow)
        img.save(out_dir / f"slide{idx:02d}.png")
        if overflow:
            report.append((idx, overflow))
    for idx, ov in report:
        print(f"slide {idx}: {len(ov)} possible text overflows -> {ov[:4]}")
    print("rendered to", out_dir)


if __name__ == "__main__":
    deck = ROOT / "IITBBS_Impact_Dissipating_Crash_Barrier_Internal_Confidential.pptx"
    only = {int(a) for a in sys.argv[1:]} or None
    render(deck, ROOT / ".build" / "preview", only)
