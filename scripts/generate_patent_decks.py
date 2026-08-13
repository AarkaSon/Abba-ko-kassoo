#!/usr/bin/env python3
"""Generate two editable patent-value PowerPoint decks and a supporting workbook.

Inputs are the inventor-supplied DOCX files in the repository. Charts are native
PowerPoint charts; all narrative graphics are editable PowerPoint shapes. The
internal deck contains tested composition details; the external deck redacts
claim-enabling formulation details and carries an IP-approval warning.
"""

from __future__ import annotations

import math
import os
import shutil
from pathlib import Path
from typing import Iterable, Sequence
from zipfile import ZipFile

from PIL import Image, ImageEnhance, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.text import MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
BUILD = ROOT / ".build_patent_deck"
MEDIA = BUILD / "article_media"
CROPS = BUILD / "crops"
ARTICLE = ROOT / "Al-Amir_Research Article_R1.docx"
DISCLOSURE = ROOT / "Annexure - V Invention Disclosure for Patents & Copyrights.docx"
QUESTIONNAIRE = ROOT / "ANNEXURE - V(B).docx"
TECH_PROFILE = ROOT / "Annexure - VI Technology Profile for Patent & Copyrights.docx"
LOGO = ROOT / "assets" / "iitbbs_logo_official.png"

SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)

# Palette
NAVY = "08243A"
NAVY_2 = "0D3651"
INK = "172733"
MUTED = "596C78"
PALE = "F4F7F9"
WHITE = "FFFFFF"
GRID = "D8E2E8"
BLUE = "2B78A6"
CYAN = "22B8C7"
TEAL = "178D86"
GREEN = "2E9D69"
ORANGE = "F39C3D"
RED = "D95D5D"
YELLOW = "F2C94C"
LIGHT_BLUE = "DDEEF5"
LIGHT_CYAN = "DDF5F4"
LIGHT_ORANGE = "FFF0DF"
LIGHT_RED = "FBE7E7"
LIGHT_GREEN = "E3F3EB"
DARK_BG = "061B2C"

FONT_HEAD = "Aptos Display"
FONT_BODY = "Aptos"

# Source data transcribed from inventor-supplied manuscript tables/charts.
MIXES = ["NWC", "LWC00", "LWC05", "LWC10", "LWC15"]
FIBRE_VOL = [0.0, 0.0, 0.5, 1.0, 1.5]
DENSITY = [2466, 1885, 1923, 1951, 1987]
SLUMP = [130, 165, 120, 112, 87]
COMP = [49.62, 51.10, 53.70, 56.80, 55.40]
TENSILE = [4.32, 3.36, 4.77, 5.92, 7.60]
PRISM_N1 = [2, 1, 5, 15, 35]
PRISM_N2 = [4, 2, 20, 45, 221]
PRISM_E1 = [26.49, 13.24, 66.22, 198.65, 463.52]
PRISM_E2 = [52.97, 26.49, 264.87, 595.96, 2927.81]
DISP = [0.28, 0.08, 0.14, 0.85, 0.71]
DISP_SD = [0.0050, 0.0163, 0.0144, 0.0173, 0.0686]
COST = [4945.24, 4411.62, 8336.62, 12261.62, 16191.62]
COST_E = [93.36, 166.54, 31.47, 20.57, 5.53]
BARRIER_MIXES = ["NWC", "LWC15"]
BARRIER_N1 = [4, 7]
BARRIER_N2 = [10, 20]
BARRIER_E1 = [549.2, 961.1]
BARRIER_E2 = [1373, 2746]
BARRIER_DEF = [4.22, 12.70]

MIX_COMPONENTS = {
    "NWC": [338.8, 87.4, 173.63, 795.62, 948.8, 0.0, 0.0, 1.90],
    "LWC00": [422.4, 105.6, 195.16, 466.46, 0.0, 596.31, 0.0, 2.79],
    "LWC05": [422.4, 105.6, 195.16, 466.46, 0.0, 596.31, 39.25, 2.79],
    "LWC10": [422.4, 105.6, 195.16, 466.46, 0.0, 596.31, 78.50, 2.79],
    "LWC15": [422.4, 105.6, 195.16, 466.46, 0.0, 596.31, 117.8, 2.79],
}
COMPONENT_NAMES = ["OPC", "GGBS", "Water", "Fine agg.", "Natural coarse", "SFA", "Steel fibre", "PCE"]

# Calculated headline metrics.
DENSITY_REDUCTION = 1 - DENSITY[-1] / DENSITY[0]
COMP_GAIN = COMP[-1] / COMP[0] - 1
TENSILE_GAIN = TENSILE[-1] / TENSILE[0] - 1
PRISM_GAIN_NWC = PRISM_E2[-1] / PRISM_E2[0]
PRISM_GAIN_LWC0 = PRISM_E2[-1] / PRISM_E2[1]
COST_ENERGY_GAIN = COST_E[0] / COST_E[-1]
BARRIER_ENERGY_GAIN = BARRIER_E2[-1] / BARRIER_E2[0]
BARRIER_DEF_GAIN = BARRIER_DEF[-1] / BARRIER_DEF[0]


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def extract_media() -> None:
    MEDIA.mkdir(parents=True, exist_ok=True)
    with ZipFile(ARTICLE) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/"):
                (MEDIA / Path(name).name).write_bytes(zf.read(name))


def crop_image(src: Path, key: str, size=(1200, 700), enhance=False) -> Path:
    CROPS.mkdir(parents=True, exist_ok=True)
    dst = CROPS / f"{key}.jpg"
    if dst.exists():
        return dst
    im = Image.open(src).convert("RGB")
    im = ImageOps.fit(im, size, method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))
    if enhance:
        im = ImageEnhance.Contrast(im).enhance(1.08)
        im = ImageEnhance.Color(im).enhance(0.92)
    im.save(dst, quality=92, optimize=True)
    return dst


def contain_image(src: Path, key: str, size=(1200, 700), background=(255, 255, 255)) -> Path:
    """Fit a technical figure without cropping labels or subpanels."""
    CROPS.mkdir(parents=True, exist_ok=True)
    dst = CROPS / f"{key}.jpg"
    if dst.exists():
        return dst
    im = Image.open(src).convert("RGB")
    fitted = ImageOps.contain(im, size, method=Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", size, background)
    canvas.paste(fitted, ((size[0] - fitted.width)//2, (size[1] - fitted.height)//2))
    canvas.save(dst, quality=94, optimize=True)
    return dst


def darken_image(src: Path, key: str, size=(1000, 1000), factor=0.55) -> Path:
    CROPS.mkdir(parents=True, exist_ok=True)
    dst = CROPS / f"{key}.jpg"
    im = Image.open(src).convert("RGB")
    im = ImageOps.fit(im, size, method=Image.Resampling.LANCZOS)
    im = ImageEnhance.Brightness(im).enhance(factor)
    im = ImageEnhance.Contrast(im).enhance(1.05)
    im.save(dst, quality=92, optimize=True)
    return dst


def set_run(run, size=12, color=INK, bold=False, font=FONT_BODY, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def add_text(slide, x, y, w, h, text, size=14, color=INK, bold=False,
             font=FONT_BODY, align=PP_ALIGN.LEFT, valign=MSO_VERTICAL_ANCHOR.TOP,
             margin=0.04, line_spacing=1.0, name=None, rotation=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    box.rotation = rotation
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    set_run(r, size, color, bold, font)
    return box


def add_rich_text(slide, x, y, w, h, parts, size=14, color=INK,
                  align=PP_ALIGN.LEFT, valign=MSO_VERTICAL_ANCHOR.TOP,
                  margin=0.04, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for part in parts:
        r = p.add_run()
        r.text = part.get("text", "")
        set_run(r, part.get("size", size), part.get("color", color),
                part.get("bold", False), part.get("font", FONT_BODY),
                part.get("italic", False))
    return box


def add_shape(slide, shape_type, x, y, w, h, fill=WHITE, line=GRID,
              radius=True, line_width=0.8, name=None):
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(line)
    shp.line.width = Pt(line_width)
    return shp


def add_rect(slide, x, y, w, h, fill=WHITE, line=GRID, radius=True,
             line_width=0.8, name=None):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    return add_shape(slide, st, x, y, w, h, fill, line, line_width=line_width, name=name)


def add_line(slide, x1, y1, x2, y2, color=GRID, width=1.0, dash=False, name=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        ln.name = name
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return ln


def add_pill(slide, x, y, w, text, fill=LIGHT_BLUE, color=BLUE, size=9.5, line=None):
    shp = add_rect(slide, x, y, w, 0.32, fill, line or fill, radius=True)
    shp.adjustments[0] = 0.5
    add_text(slide, x + 0.04, y + 0.015, w - 0.08, 0.27, text, size, color, True,
             align=PP_ALIGN.CENTER, valign=MSO_VERTICAL_ANCHOR.MIDDLE, margin=0)
    return shp


def add_card(slide, x, y, w, h, title, body="", accent=BLUE, fill=WHITE,
             title_size=13, body_size=10.5, icon=None):
    add_rect(slide, x, y, w, h, fill, GRID, radius=True, line_width=0.7)
    add_rect(slide, x, y, 0.07, h, accent, accent, radius=False, line_width=0)
    tx = x + 0.23
    if icon:
        add_text(slide, tx, y + 0.20, 0.35, 0.35, icon, 17, accent, True,
                 align=PP_ALIGN.CENTER, valign=MSO_VERTICAL_ANCHOR.MIDDLE, margin=0)
        tx += 0.46
    add_text(slide, tx, y + 0.18, w - (tx - x) - 0.18, 0.35, title, title_size, INK, True,
             FONT_HEAD, valign=MSO_VERTICAL_ANCHOR.MIDDLE, margin=0)
    if body:
        add_text(slide, x + 0.23, y + 0.62, w - 0.43, h - 0.76, body, body_size, MUTED,
                 False, FONT_BODY, margin=0, line_spacing=1.05)


def add_metric(slide, x, y, w, h, value, label, accent=CYAN, fill=WHITE,
               value_size=25, sub=None):
    add_rect(slide, x, y, w, h, fill, GRID, radius=True, line_width=0.7)
    add_text(slide, x + 0.18, y + 0.16, w - 0.36, 0.48, value, value_size, accent, True,
             FONT_HEAD, margin=0, valign=MSO_VERTICAL_ANCHOR.MIDDLE)
    add_text(slide, x + 0.18, y + 0.70, w - 0.36, 0.44, label, 10.5, INK, True,
             margin=0)
    if sub:
        add_text(slide, x + 0.18, y + 1.12, w - 0.36, h - 1.24, sub, 8.5, MUTED,
                 margin=0)


def add_logo(slide, dark=False):
    # A small white plaque keeps the official mark legible on every slide.
    add_rect(slide, 12.22, 0.14, 0.82, 0.64, WHITE, WHITE, radius=True, line_width=0)
    slide.shapes.add_picture(str(LOGO), Inches(12.34), Inches(0.18), width=Inches(0.58), height=Inches(0.54))


def add_header(slide, title, section, slide_no, status, dark=False):
    bg = DARK_BG if dark else PALE
    txt = WHITE if dark else INK
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    add_text(slide, 0.70, 0.18, 2.75, 0.22, section.upper(), 8.5,
             CYAN if dark else BLUE, True, margin=0)
    # Reserve two lines for research-style titles; this prevents descenders from
    # colliding with the divider on longer headings.
    title_size = 22.5 if len(title) <= 58 else 20.8
    add_text(slide, 0.70, 0.37, 10.85, 0.67, title, title_size, txt, True, FONT_HEAD,
             margin=0, valign=MSO_VERTICAL_ANCHOR.MIDDLE, line_spacing=0.92)
    add_logo(slide, dark)
    add_line(slide, 0.70, 1.10, 12.62, 1.10, "284A60" if dark else GRID, 0.8)
    add_footer(slide, slide_no, status, dark)


def add_footer(slide, slide_no, status, dark=False):
    c = "AEC0CA" if dark else MUTED
    add_line(slide, 0.70, 7.08, 12.62, 7.08, "284A60" if dark else GRID, 0.7)
    add_text(slide, 0.72, 7.16, 10.7, 0.18, status, 7.6, c, True, margin=0)
    add_text(slide, 12.00, 7.14, 0.60, 0.19, f"{slide_no:02d}", 8.3, c, True,
             align=PP_ALIGN.RIGHT, margin=0)


def add_source(slide, x, y, w, text, dark=False, size=7.0):
    color = "A9BBC4" if dark else MUTED
    add_text(slide, x, y, w, 0.28, "Source: " + text, size, color, False,
             margin=0, valign=MSO_VERTICAL_ANCHOR.MIDDLE)


def add_notes(slide, text: str):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    tf.text = text.strip()
    for p in tf.paragraphs:
        for r in p.runs:
            set_run(r, 12, INK, False, FONT_BODY)


def add_picture_cover(slide, src: Path, x, y, w, h, key, enhance=True):
    ratio = max(1, round(w * 170)), max(1, round(h * 170))
    cropped = crop_image(src, key, ratio, enhance)
    return slide.shapes.add_picture(str(cropped), Inches(x), Inches(y), Inches(w), Inches(h))


def add_picture_contain(slide, src: Path, x, y, w, h, key):
    ratio = max(1, round(w * 190)), max(1, round(h * 190))
    fitted = contain_image(src, key, ratio)
    return slide.shapes.add_picture(str(fitted), Inches(x), Inches(y), Inches(w), Inches(h))


def style_chart(chart, font_size=8.5, legend=False, legend_pos=XL_LEGEND_POSITION.BOTTOM):
    chart.font.name = FONT_BODY
    chart.font.size = Pt(font_size)
    chart.font.color.rgb = rgb(MUTED)
    chart.has_legend = legend
    if legend:
        chart.legend.position = legend_pos
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT_BODY
        chart.legend.font.size = Pt(8)
        chart.legend.font.color.rgb = rgb(MUTED)
    # python-pptx does not currently expose chart-area/plot-area formatting.
    # PowerPoint's default white chart canvas is retained.


def color_series(series, color):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(color)
    series.format.line.color.rgb = rgb(color)


def color_points(series, colors: Sequence[str]):
    for point, color in zip(series.points, colors):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(color)
        point.format.line.color.rgb = rgb(color)


def add_column_chart(slide, x, y, w, h, categories, series_data,
                     colors, y_max=None, number_format="0.0", legend=False,
                     data_labels=True, y_title=None, gap=55, title=None):
    data = ChartData()
    data.categories = categories
    for s_name, vals in series_data:
        data.add_series(s_name, vals)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                   Inches(w), Inches(h), data)
    chart = frame.chart
    style_chart(chart, 8.0, legend)
    for i, s in enumerate(chart.series):
        color_series(s, colors[i % len(colors)])
    chart.plots[0].gap_width = gap
    chart.category_axis.tick_labels.font.name = FONT_BODY
    chart.category_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.format.line.color.rgb = rgb(GRID)
    chart.value_axis.tick_labels.font.name = FONT_BODY
    chart.value_axis.tick_labels.font.size = Pt(7.5)
    chart.value_axis.format.line.color.rgb = rgb(GRID)
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb(GRID)
    chart.value_axis.minimum_scale = 0
    if y_max:
        chart.value_axis.maximum_scale = y_max
    if y_title:
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = y_title
        chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    if data_labels:
        plot = chart.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
        labels.show_value = True
        labels.number_format = number_format
        labels.font.name = FONT_BODY
        labels.font.size = Pt(7.5)
        labels.font.color.rgb = rgb(INK)
    if title:
        add_text(slide, x + 0.05, y + 0.02, w - 0.1, 0.26, title, 10.5, INK, True,
                 margin=0, align=PP_ALIGN.CENTER)
    return chart


def add_bar_chart(slide, x, y, w, h, categories, values, color=BLUE,
                  number_format="0", y_max=None, title=None, point_colors=None):
    data = ChartData()
    data.categories = categories
    data.add_series("Value", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y),
                                   Inches(w), Inches(h), data)
    chart = frame.chart
    style_chart(chart, 8.2, False)
    s = chart.series[0]
    color_series(s, color)
    if point_colors:
        color_points(s, point_colors)
    chart.plots[0].gap_width = 58
    chart.category_axis.reverse_order = True
    chart.category_axis.tick_labels.font.name = FONT_BODY
    chart.category_axis.tick_labels.font.size = Pt(8.3)
    chart.category_axis.format.line.color.rgb = rgb(GRID)
    chart.value_axis.minimum_scale = 0
    if y_max:
        chart.value_axis.maximum_scale = y_max
    chart.value_axis.tick_labels.font.size = Pt(7.5)
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb(GRID)
    chart.value_axis.format.line.color.rgb = rgb(GRID)
    chart.plots[0].has_data_labels = True
    labels = chart.plots[0].data_labels
    labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    labels.show_value = True
    labels.number_format = number_format
    labels.font.name = FONT_BODY
    labels.font.size = Pt(8)
    labels.font.color.rgb = rgb(INK)
    if title:
        add_text(slide, x + 0.05, y + 0.02, w - 0.1, 0.26, title, 10.5, INK, True,
                 margin=0, align=PP_ALIGN.CENTER)
    return chart


def add_doughnut(slide, x, y, w, h, categories, values, colors):
    data = ChartData()
    data.categories = categories
    data.add_series("Share", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y),
                                   Inches(w), Inches(h), data)
    chart = frame.chart
    style_chart(chart, 8.0, True, XL_LEGEND_POSITION.BOTTOM)
    chart.plots[0].doughnut_hole_size = 66
    color_points(chart.series[0], colors)
    chart.plots[0].has_data_labels = False
    return chart


def table_text_style(cell, size=8.5, color=INK, bold=False, align=PP_ALIGN.CENTER):
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            set_run(r, size, color, bold, FONT_BODY)


def add_table(slide, x, y, w, h, rows: Sequence[Sequence], col_widths=None,
              header_fill=NAVY, header_color=WHITE, body_fill=WHITE,
              font_size=8.5, first_col_left=False, highlight_rows=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(w * cw / total)
    highlight_rows = highlight_rows or {}
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = str(val)
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
                color = header_color
                bold = True
            else:
                cell.fill.fore_color.rgb = rgb(highlight_rows.get(ri, body_fill if ri % 2 else "F8FAFB"))
                color = INK
                bold = ci == 0
            cell.border if False else None
            table_text_style(cell, font_size if ri else font_size - 0.2, color, bold,
                             PP_ALIGN.LEFT if (first_col_left and ci == 0) else PP_ALIGN.CENTER)
    return table


def add_chevron(slide, x, y, w, h, fill=BLUE):
    shp = add_shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, fill, fill, line_width=0)
    return shp


def add_circle_label(slide, x, y, d, text, fill=BLUE, color=WHITE, size=11):
    add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill, fill, line_width=0)
    add_text(slide, x, y, d, d, text, size, color, True,
             align=PP_ALIGN.CENTER, valign=MSO_VERTICAL_ANCHOR.MIDDLE, margin=0)


def add_bullets(slide, x, y, w, h, bullets, size=12, color=INK,
                bullet_color=CYAN, gap=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ""
        p.space_after = Pt(gap * 72)
        r1 = p.add_run(); r1.text = "●  "; set_run(r1, size - 1, bullet_color, True)
        if isinstance(item, tuple):
            lead, rest = item
            r2 = p.add_run(); r2.text = lead; set_run(r2, size, color, True)
            r3 = p.add_run(); r3.text = rest; set_run(r3, size, color, False)
        else:
            r2 = p.add_run(); r2.text = str(item); set_run(r2, size, color, False)
    return box


def base_prs(title, subject, comments) -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    cp = prs.core_properties
    cp.title = title
    cp.subject = subject
    cp.author = "IIT Bhubaneswar invention team"
    cp.keywords = "crash barrier, lightweight concrete, impact energy, patent, IIT Bhubaneswar"
    cp.comments = comments
    cp.company = "Indian Institute of Technology Bhubaneswar"
    return prs


def status_text(internal: bool) -> str:
    return ("CONFIDENTIAL • PRE-FILING IITBBS IPR COMMITTEE EVALUATION • NOT YET FILED"
            if internal else
            "EXTERNAL REDACTED DRAFT • NOT FOR DISTRIBUTION BEFORE FILING • IITBBS IPR APPROVAL REQUIRED")


def title_slide(prs, internal: bool):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(DARK_BG)
    status = status_text(internal)
    # Right-side experimental image.
    hero = darken_image(MEDIA / "image8.jpeg", "title_test_rig", (1000, 1000), 0.62)
    slide.shapes.add_picture(str(hero), Inches(8.75), Inches(0), width=Inches(4.58), height=Inches(7.5))
    add_rect(slide, 8.70, 0, 0.08, 7.5, CYAN, CYAN, radius=False, line_width=0)
    add_logo(slide, True)
    add_pill(slide, 0.78, 0.62, 2.70,
             "IITBBS IPR COMMITTEE REVIEW" if internal else "EXTERNAL TECHNOLOGY OVERVIEW",
             "14364B", CYAN, 9.0, "14364B")
    title = ("Novel concrete composition for\nimpact-energy-dissipating\ncrash barriers")
    add_text(slide, 0.78, 1.25, 7.45, 2.05, title, 30.5, WHITE, True, FONT_HEAD,
             margin=0, line_spacing=0.88)
    sub = ("Steel-fibre reinforced sintered-fly-ash lightweight aggregate concrete"
           if internal else
           "A lightweight, ductile concrete platform for safer and more sustainable rigid barriers")
    add_text(slide, 0.82, 3.57, 7.35, 0.70, sub, 15.0, "C6D8E2", False,
             margin=0, line_spacing=1.05)
    add_line(slide, 0.82, 4.54, 7.75, 4.54, "315267", 1.0)
    add_text(slide, 0.82, 4.82, 7.40, 0.60,
             "Haruna Al-Amir Saleh  •  Pratik Kanungo  •  Anush K. Chandrappa  •  Dinakar Pasla",
             11.2, WHITE, True, margin=0)
    add_text(slide, 0.82, 5.48, 5.8, 0.48,
             "School of Infrastructure  |  IIT Bhubaneswar  |  13 August 2026",
             10.2, "AFC3CE", margin=0)
    add_text(slide, 0.82, 6.27, 7.40, 0.48,
             "Current status: not yet filed  •  Purpose: IPR approval to proceed with patent filing",
             9.5, ORANGE, True, margin=0)
    add_text(slide, 9.07, 6.53, 3.85, 0.42,
             "Source: inventor-supplied manuscript Fig. 7 — 1:3 barrier impact setup",
             7.1, "D8E5EB", False, align=PP_ALIGN.RIGHT, margin=0)
    add_footer(slide, 1, status, True)
    add_notes(slide, f"""
OPENING — approximately 45 seconds

• This presentation evaluates an experimentally reduced-to-practice material platform for impact-energy-dissipating crash barriers. It is not a legal patentability opinion.
• The invention team is Haruna Al-Amir Saleh, Pratik Kanungo, Dr Anush K. Chandrappa and Prof Dinakar Pasla, School of Infrastructure, IIT Bhubaneswar.
• The inventors confirm that the patent has not yet been filed. This is the pre-filing presentation to the IIT Bhubaneswar IPR section seeking approval to proceed with professional drafting and filing. A related manuscript was submitted on 18 July 2026, so confidentiality and disclosure timing are critical. Treat this deck accordingly: {status}.
• The central question is not whether the concrete is merely lighter. It is whether the tested combination creates a defensible, multi-mechanism improvement in energy absorption, post-cracking integrity and controlled damage while retaining structural-grade strength.
• The presentation separates measured evidence from proposed occupant-safety benefit. Full-scale vehicle testing is still required before any claim of reduced occupant injury or standards compliance.
""")


def slide_executive(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "The value case in one view", "Executive thesis", 2, status)
    add_text(slide, 0.73, 1.17, 11.85, 0.44,
             "A rigid-barrier material that trades brittle force transfer for controlled, multi-scale energy dissipation.",
             16.2, NAVY, True, FONT_HEAD, margin=0)
    metrics = [
        (f"{DENSITY_REDUCTION*100:.1f}%", "lower reported density", "LWC15 vs NWC", BLUE),
        (f"{TENSILE_GAIN*100:.0f}%", "higher splitting tensile strength", "7.60 vs 4.32 MPa", TEAL),
        (f"{PRISM_GAIN_NWC:.1f}×", "prism failure-energy capacity", "2,927.8 vs 53.0 N·m", ORANGE),
        (f"{BARRIER_ENERGY_GAIN:.1f}×", "scaled-barrier failure energy", "2,746 vs 1,373 N·m", RED),
    ]
    for i, (v, lab, sub, col) in enumerate(metrics):
        add_metric(slide, 0.73 + i*3.02, 1.82, 2.78, 1.62, v, lab, col, WHITE, 24, sub)
    add_card(slide, 0.73, 3.78, 3.80, 2.70, "Protectable technical thesis",
             "Specific material architecture + dual dissipation mechanism + crash-barrier use + tested performance envelope. Claims should centre the integrated system—not isolated known constituents.",
             CYAN, WHITE, 13.5, 10.7, "01")
    add_card(slide, 4.75, 3.78, 3.80, 2.70, "Commercial value pathway",
             "Precast-compatible manufacturing; lower barrier self-weight; improved damage control; potential transport, deck-load and lifecycle advantages. These benefits require pilot-scale TEA/LCA validation.",
             GREEN, WHITE, 13.5, 10.7, "02")
    add_card(slide, 8.77, 3.78, 3.80, 2.70, "Evidence boundary",
             "Material and repeated-impact results are strong signals. The 1:3 barrier test used one specimen per mix and did not measure ASI/THIV, vehicle containment or redirection. Occupant benefit remains a testable hypothesis.",
             RED, WHITE, 13.5, 10.7, "03")
    add_source(slide, 0.73, 6.67, 11.7,
               "Inventor dataset, Al-Amir et al. manuscript (submitted 18 Jul 2026), Tables/Figs 8–12; calculations in accompanying workbook.")
    patent_thesis_note = ("The patent thesis should not be “steel fibre is useful” or “lightweight aggregate is useful”; both are known. The potentially protectable value is the integrated composition, the use of SFA and hooked-end fibres in a crash-barrier system, the dual energy-dissipation mechanism, and the demonstrated performance combination."
                          if internal else
                          "The potentially protectable value is the integrated material architecture, crash-barrier application, dual energy-dissipation mechanism and demonstrated performance combination. Exact formulation and processing details are redacted from this version.")
    add_notes(slide, f"""
KEY MESSAGE — approximately 60 seconds

• Lead with the evidence: relative to normal-weight concrete, the selected lightweight mixture has 19.4% lower reported density and 75.9% higher splitting tensile strength. In the repeated drop-weight prism test it absorbed 55.3 times the failure energy of NWC.
• At the barrier scale, the one-third model absorbed twice the cumulative pendulum energy before failure and deformed approximately three times as much as the NWC model. That is the most application-relevant result, but it is preliminary because only one barrier per mix was tested.
• {patent_thesis_note}
• Commercially, the material can use conventional concrete/precast infrastructure. The cost per unit of laboratory impact energy is favourable, but no lifecycle cost analysis or environmental life-cycle assessment has yet been performed.
• Do not state that occupant injury has been reduced. Say that the measured deformation and energy absorption support a hypothesis of lower force transmission that must be tested at full scale.
""")


def slide_global(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "The need is global: safety, materials and circularity intersect", "Global evidence", 3, status)
    add_doughnut(slide, 0.70, 1.31, 4.10, 3.25,
                 ["Upper-middle", "Lower-middle", "Low", "High"], [35, 44, 13, 8],
                 [BLUE, CYAN, ORANGE, GRID])
    add_text(slide, 2.02, 2.26, 1.48, 0.78, "92%", 27, NAVY, True, FONT_HEAD,
             align=PP_ALIGN.CENTER, valign=MSO_VERTICAL_ANCHOR.MIDDLE, margin=0)
    add_text(slide, 1.90, 2.95, 1.72, 0.40, "of deaths in LMICs", 9.5, MUTED, True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, 0.84, 4.59, 3.85, 0.46, "Global road-traffic deaths by country income group, 2021", 10.5, INK, True,
             align=PP_ALIGN.CENTER, margin=0)
    # Global context cards.
    cards = [
        ("1.19 M", "road deaths/year", "WHO 2023; 2021 estimate", RED),
        ("50 Gt", "sand + gravel/year", "UNEP 2022 estimate", ORANGE),
        ("4.0 Gt", "cement produced, 2024", "USGS 2025 • ~0.6 tCO₂/t (IEA)", BLUE),
        (">1 Gt", "coal fly ash/year", "IEA CCC 2020; global order", TEAL),
    ]
    positions = [(5.12, 1.33), (8.87, 1.33), (5.12, 3.16), (8.87, 3.16)]
    for (v, lab, sub, col), (x, y) in zip(cards, positions):
        add_metric(slide, x, y, 3.42, 1.54, v, lab, col, WHITE, 23, sub)
    add_rect(slide, 5.12, 5.10, 7.17, 1.05, NAVY, NAVY, radius=True, line_width=0)
    add_text(slide, 5.38, 5.30, 6.65, 0.60,
             "Opportunity: reduce crash severity while substituting a portion of virgin aggregate and clinker-intensive binder inputs.\nBoundary: this project has not yet completed full-scale crash testing or a product LCA.",
             11.0, WHITE, True, margin=0, line_spacing=1.0)
    add_source(slide, 0.73, 6.45, 11.9,
               "WHO, Global Status Report on Road Safety 2023, pp. 4, 14; UNEP, Sand and Sustainability (2022); USGS, MCS—Cement (2025); IEA, Cement (2023); Reid, Carpenter & Masili, IEA CCC/303 (2020).")
    lca_boundary = ("The mix uses GGBS and sintered fly ash aggregate, but sintering consumes energy and the high steel-fibre dosage has embodied impacts."
                    if internal else
                    "The technology incorporates industrial by-product pathways, but aggregate processing and the reinforcing phase also carry embodied impacts; exact constituents are redacted.")
    add_notes(slide, f"""
GLOBAL CONTEXT — approximately 75 seconds

• The global road-safety burden is about 1.19 million deaths annually. WHO estimates that 92% occur in low- and middle-income countries. The pie chart is a global dataset: 35% upper-middle, 44% lower-middle, 13% low-income and 8% high-income countries.
• At the same time, construction operates at planetary material scale: UNEP estimates 50 billion tonnes of sand and gravel are used each year; USGS estimates about 4.0 billion tonnes of cement were produced in 2024, while IEA reports direct cement-emissions intensity just under 0.6 tonnes of CO₂ per tonne.
• Coal combustion also generates more than one billion tonnes of fly ash annually by order of magnitude. Availability and specification vary regionally, so this figure supports the resource opportunity—not a guaranteed feedstock claim.
• The invention sits at the intersection of three needs: safer roadside infrastructure, lower self-weight and virgin-aggregate demand, and productive use of industrial by-products.
• Be explicit about the environmental boundary. {lca_boundary} A comparative, region-specific LCA is required before claiming a net carbon reduction.
• Transition: global need alone does not establish novelty. The next slide maps the prior-art trade space.
""")


def slide_prior_art(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Prior art leaves an unresolved multi-objective gap", "Landscape", 4, status)
    lightweight_gap = "No SFA–fibre proof set" if internal else "No integrated proof set"
    rows = [
        ["Approach", "Containment / strength", "Energy dissipation", "Damage control", "Scale / cost", "Residual gap"],
        ["Rigid NWC barrier", "High", "Low–moderate", "Brittle / debris risk", "Mature", "Harsh impact pulse"],
        ["Steel / cable", "System-dependent", "High via deflection", "Repair after impact", "Corrosion / clear zone", "Space + maintenance"],
        ["Rubberised + fibre", "Strength trade-off", "Improved", "Improved", "Processing complexity", "Durability / scale"],
        ["UHPC / ECC", "Very high", "Mix-dependent", "Good crack control", "High material cost", "Stiffness / economics"],
        ["Lightweight barrier concepts", "Demonstrated in prior art", "Potential", "Design-specific", "Limited adoption", lightweight_gap],
        ["This invention", "M40+ in tested mixes", "High in repeated impact", "Controlled cracking observed", "Precast-compatible", "Full-scale proof pending"],
    ]
    add_table(slide, 0.72, 1.35, 11.90, 3.78, rows, [1.7, 1.35, 1.35, 1.45, 1.45, 2.10],
              font_size=8.6, first_col_left=True, highlight_rows={6: LIGHT_CYAN})
    add_rect(slide, 0.73, 5.44, 11.88, 0.95, WHITE, GRID, radius=True)
    add_text(slide, 0.98, 5.63, 2.10, 0.30, "PATENT POSITIONING", 9.0, BLUE, True, margin=0)
    positioning = ("Avoid a broad “first lightweight concrete barrier” assertion: lightweight-concrete barrier concepts appear in conference literature. Centre novelty on the specific integrated composition, SFA function, hooked-fibre crack control, quantified synergy and manufacturing/control envelope."
                   if internal else
                   "Avoid a broad “first lightweight concrete barrier” assertion: lightweight-concrete barrier concepts appear in conference literature. Centre differentiation on the integrated architecture, measured synergy, controlled-damage response and protected manufacturing envelope.")
    add_text(slide, 2.98, 5.55, 9.28, 0.58, positioning,
             10.6, INK, True, margin=0, line_spacing=1.0)
    add_source(slide, 0.73, 6.55, 11.85,
               "Zain & Mohammed (2015); Raj et al. (2020), ASCE J. Mater. Civ. Eng.; Shen et al. (2014), Constr. Build. Mater.; Bonin et al. (2004), II Int’l Congress SIIV proceedings; Grzebieta et al. (2005), 19th ESV Conference proceedings; Xue et al. (2025).")
    exact_thesis = ("The narrower and more defensible thesis is the integrated OPC/GGBS/SFA/hooked-end-fibre composition, its use as the load-bearing crash-barrier material, and the observed synergy: structural strength, lower reported density, major post-cracking impact resistance and controlled fragmentation."
                    if internal else
                    "The narrower differentiation thesis is the protected integrated architecture, its crash-barrier use and the measured combination of structural strength, lower reported density, post-cracking impact resistance and controlled damage. Exact formulation details are intentionally omitted from this version.")
    search_scope = ("including composition ranges, SFA manufacture/specification, fibre geometry and barrier-system claims"
                    if internal else
                    "including material-platform, manufacturing, barrier-article and system-use claims")
    add_notes(slide, f"""
PRIOR ART AND DIFFERENTIATION — approximately 90 seconds

• Conventional rigid concrete barriers are mature and effective for containment, but their low deflection can create a severe vehicle pulse, and brittle damage can generate debris under severe impact.
• Flexible and semi-rigid systems dissipate energy through global deflection, but require working width and post-impact repair. Rubberised, UHPC and ECC concepts each solve part of the problem, while introducing strength, cost, stiffness or processing trade-offs.
• A critical patentability point: prior conference literature includes lightweight concrete road-safety barrier concepts—for example Bonin and co-authors at the 2004 SIIV congress. Therefore, do not rely on a sweeping “first lightweight barrier” proposition.
• {exact_thesis}
• The literature and patent searches documented by the inventors are preliminary. A professional claim-by-claim patentability and freedom-to-operate search is still required, {search_scope}.
""")


def slide_invention(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Invention architecture: three functions, one controlled response", "Technical concept", 5, status)
    if internal:
        items = [
            ("SFA lightweight aggregate", "Porous particle crushing + internal damping", BLUE, "A"),
            ("Hooked-end steel fibre", "Crack bridging + debonding + pull-out", ORANGE, "B"),
            ("OPC + 20% GGBS matrix", "M40+ strength + fibre/aggregate load transfer", TEAL, "C"),
        ]
    else:
        items = [
            ("Porous lightweight aggregate", "Local crushing + internal damping", BLUE, "A"),
            ("Discrete crack-bridging reinforcement", "Debonding + pull-out + integrity", ORANGE, "B"),
            ("Lower-clinker structural matrix", "Strength + interfacial load transfer", TEAL, "C"),
        ]
    for i, (title, body, col, lab) in enumerate(items):
        y = 1.39 + i*1.34
        add_circle_label(slide, 0.82, y + 0.12, 0.54, lab, col, WHITE, 12)
        add_card(slide, 1.18, y, 4.45, 1.04, title, body, col, WHITE, 13.0, 10.0)
        add_chevron(slide, 5.85, y + 0.30, 0.62, 0.42, col)
    # Output block.
    add_rect(slide, 6.66, 1.35, 5.70, 3.82, NAVY, NAVY, radius=True, line_width=0)
    add_text(slide, 7.00, 1.67, 5.08, 0.38, "DESIGNED IMPACT RESPONSE", 10.0, CYAN, True,
             margin=0)
    outputs = [
        ("01", "Crack initiation", "Fibres redistribute tensile stress and delay localisation."),
        ("02", "Controlled deformation", "Aggregate crushing and fibre pull-out consume energy."),
        ("03", "Residual integrity", "Bridging restrains large fragments and preserves load path."),
        ("04", "Lower transmitted peak force", "Mechanistic hypothesis; requires vehicle-level validation."),
    ]
    for i, (n, t, b) in enumerate(outputs):
        y = 2.18 + i*0.70
        add_text(slide, 7.00, y, 0.42, 0.32, n, 9.0, CYAN, True, margin=0)
        add_text(slide, 7.47, y - 0.02, 1.62, 0.30, t, 10.5, WHITE, True, margin=0)
        add_text(slide, 9.16, y - 0.02, 2.80, 0.43, b, 8.8, "C4D4DC", margin=0)
    add_rect(slide, 0.82, 5.48, 11.54, 0.78, LIGHT_CYAN, LIGHT_CYAN, radius=True, line_width=0)
    thesis = ("Proposed inventive step: an intentionally semi-energy-absorbing rigid barrier in which the material—not only the barrier geometry—dissipates collision energy while retaining structural-grade strength."
              if internal else
              "Technology thesis: a semi-energy-absorbing rigid barrier that shifts part of the collision response from brittle fracture to controlled internal dissipation.")
    add_text(slide, 1.07, 5.68, 11.05, 0.38, thesis, 11.3, NAVY, True, margin=0)
    add_source(slide, 0.82, 6.47, 11.55,
               "Inventor disclosure (16 Jul 2026), §§Description/Novelty/Inventiveness; Wang & Wang (2013); Nadesan & Dinakar (2017); Sahoo et al. (2020). Mechanism diagram is an author synthesis.")
    mechanism_detail = ("First, the porous SFA can crush locally and contribute internal damping. Second, hooked-end steel fibres bridge cracks and dissipate energy through progressive debonding, friction and pull-out. Third, the cement–GGBS matrix provides the structural-grade load path and transfers stress into the aggregate and fibres."
                        if internal else
                        "First, a porous lightweight phase can crush locally and contribute internal damping. Second, a discrete crack-bridging phase dissipates energy through progressive debonding, friction and pull-out. Third, a structural matrix maintains the load path. Exact constituent identity and proportions are redacted.")
    claim_advice = ("For claim drafting, consider separate families covering composition ranges, mixing/conditioning controls, a barrier article made from the composition, and functionally linked performance features where legally supportable."
                    if internal else
                    "This redacted version should not be used to disclose formulation ranges or process controls. Release detailed claim-enabling information only after filing and IITBBS IP approval.")
    add_notes(slide, f"""
MECHANISM — approximately 80 seconds

• The concept intentionally combines three functions. {mechanism_detail}
• The expected sequence is crack initiation, distributed bridging and controlled deformation, followed by progressive internal energy dissipation and improved residual integrity.
• This mechanism is consistent with published composite-material literature, but the invention’s value proposition is the protected integration and crash-barrier application supported by the experimental results.
• Phrase the final outcome carefully. The experiments demonstrate increased specimen deformation and cumulative absorbed energy. They do not directly demonstrate a lower peak vehicle force, reduced ASI or reduced occupant injury. Those remain mechanistically plausible validation targets.
• {claim_advice}
""")


def slide_program(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Evidence ladder: from constituents to a scaled barrier", "Experimental programme", 6, status)
    constituent_desc = ("XRF, physical properties, SFA absorption, fibre geometry" if internal
                        else "Chemistry, physical properties and reinforcement quality")
    screening_desc = ("NWC + 4 LWC variants; 0–1.5% fibre by volume" if internal
                      else "NWC + 4 lightweight variants; increasing reinforcement")
    steps = [
        ("01", "Constituent\ncharacterisation", constituent_desc, BLUE),
        ("02", "Mixture\nscreening", screening_desc, CYAN),
        ("03", "Static\nperformance", "Reported density, slump, 28 d compression and split tension", TEAL),
        ("04", "Repeated\nimpact", "13.5 kg drop hammer; N₁/N₂, energy, acceleration → displacement", ORANGE),
        ("05", "Barrier-scale\nvalidation", "1:3 New Jersey models; 40 kg pendulum; NWC vs LWC15", RED),
    ]
    for i, (n, title, body, col) in enumerate(steps):
        x = 0.70 + i*2.48
        add_circle_label(slide, x + 0.78, 1.42, 0.54, n, col, WHITE, 10)
        if i < 4:
            add_chevron(slide, x + 2.12, 1.51, 0.33, 0.34, GRID)
        add_rect(slide, x, 2.12, 2.18, 2.34, WHITE, GRID, radius=True)
        add_text(slide, x + 0.15, 2.33, 1.88, 0.68, title, 13.0, INK, True, FONT_HEAD,
                 align=PP_ALIGN.CENTER, margin=0)
        add_line(slide, x + 0.25, 3.08, x + 1.93, 3.08, col, 2.1)
        add_text(slide, x + 0.18, 3.30, 1.82, 0.90, body, 9.4, MUTED, False,
                 align=PP_ALIGN.CENTER, margin=0)
    add_rect(slide, 0.82, 4.88, 11.60, 1.15, WHITE, GRID, radius=True)
    add_text(slide, 1.03, 5.12, 2.25, 0.30, "REPLICATION / UNCERTAINTY", 9.0, RED, True, margin=0)
    add_text(slide, 3.18, 5.02, 8.95, 0.70,
             "The methods report 3 cubes + 3 cylinders per mix and 4 beams cast per mix; impact tables do not report dispersion for N₁/N₂. Barrier validation used one specimen per material. Treat barrier-scale effect sizes as preliminary, not population estimates.",
             10.7, INK, True, margin=0)
    add_source(slide, 0.82, 6.35, 11.65,
               "Al-Amir et al. manuscript, Experimental Program and Testing Methods; inventor disclosure, Experimental Data Analysis. Sample-count interpretation is based on the stated casting plan.")
    screening_note = ("Four lightweight mixtures were evaluated with fibre volume increasing from zero to 1.5%, alongside an M40 normal-weight reference."
                      if internal else
                      "Four lightweight variants with progressively increasing crack-bridging reinforcement were evaluated alongside an M40 normal-weight reference; exact proportions are redacted.")
    add_notes(slide, f"""
EXPERIMENTAL LOGIC — approximately 75 seconds

• The evidence was built in five layers: constituent characterization, mixture screening, static mechanical tests, repeated drop-weight impact testing and a one-third-scale barrier comparison.
• {screening_note}
• Static tests provide the strength and density boundary. The repeated impact test provides the key post-cracking energy evidence. Acceleration signals were processed to derive displacement. The selected highest-performing mixture was then used in a New Jersey barrier model and compared with NWC.
• This is a coherent reduction-to-practice chain and is valuable for enablement. However, the uncertainty reporting is incomplete. The article states that three cubes, three cylinders and four beams were cast per mix, but the impact result tables do not provide N1/N2 dispersion. Only one NWC and one selected lightweight barrier were cast.
• The deck therefore uses measured ratios as experimental observations—not as statistically generalized performance guarantees.
""")


def slide_mix(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Tested material design: a fibre-content performance ladder", "Tested embodiments", 7, status)
    if internal:
        rows = [["Mix", "OPC", "GGBS", "Water", "Fine agg.", "Nat. coarse", "SFA", "Steel fibre", "PCE"]]
        for mix in MIXES:
            vals = MIX_COMPONENTS[mix]
            rows.append([mix] + [f"{v:.1f}" if v else "—" for v in vals])
        add_table(slide, 0.72, 1.33, 8.80, 3.58, rows,
                  [0.75, 0.75, 0.70, 0.78, 0.88, 1.02, 0.78, 0.92, 0.62],
                  font_size=8.0, first_col_left=True, highlight_rows={5: LIGHT_CYAN})
        add_text(slide, 0.82, 4.98, 8.55, 0.32,
                 "All constituent quantities in kg/m³. Extra-water quantity intentionally omitted pending method correction.",
                 8.2, RED, True, margin=0)
        right_title = "SELECTED EMBODIMENT • LWC15"
        body = "• 1.5% fibre by concrete volume\n• 20% GGBS replacement of total binder\n• 100% SFA substitution for natural coarse aggregate\n• Hooked fibre: 30 mm × 0.5 mm; aspect ratio 60\n• Reported density: 1,987 kg/m³\n• 28 d compressive strength: 55.4 MPa"
    else:
        rows = [
            ["Programme", "Reference", "Lightweight controls", "Reinforced variants", "Selected"],
            ["Design", "M40 NWC", "Porous-LWA concrete", "Three increasing reinforcement levels", "Highest tested level"],
            ["Purpose", "Benchmark", "Separate LWA effect", "Map performance response", "Barrier model"],
        ]
        add_table(slide, 0.72, 1.45, 8.80, 2.80, rows, [1.05, 1.15, 1.55, 1.95, 1.30],
                  font_size=9.0, first_col_left=True, highlight_rows={2: LIGHT_CYAN})
        add_rect(slide, 0.75, 4.60, 8.75, 0.92, LIGHT_RED, LIGHT_RED, radius=True, line_width=0)
        add_text(slide, 0.98, 4.80, 8.30, 0.45,
                 "Formulation proportions, moisture-conditioning protocol and constituent specifications are redacted. Release only after patent filing and IITBBS IP approval.",
                 10.0, RED, True, margin=0)
        right_title = "SELECTED TECHNOLOGY EMBODIMENT"
        body = "• Structural lightweight design target\n• Industrial by-product aggregate pathway\n• Lower-clinker binder pathway\n• Discrete crack-bridging reinforcement\n• Reported density below 2,000 kg/m³\n• 28 d strength above M40 requirement"
    add_rect(slide, 9.78, 1.34, 2.60, 4.72, NAVY, NAVY, radius=True, line_width=0)
    add_text(slide, 10.02, 1.67, 2.14, 0.66, right_title, 10.6, CYAN, True, margin=0)
    add_text(slide, 10.02, 2.54, 2.10, 2.95, body, 10.0, WHITE, False, margin=0, line_spacing=1.1)
    add_pill(slide, 10.02, 5.45, 2.04, "SCALed barrier mix", "173D55", CYAN, 8.8, "173D55")
    add_source(slide, 0.74, 6.34, 11.65,
               "Al-Amir et al. manuscript, Tables 4–6 and chart caches. Density terminology retained as “reported density”; the source alternates between wet and dry density. Exact extra-water protocol omitted by user instruction.")
    if internal:
        mix_notes = """
TESTED EMBODIMENTS — approximately 70 seconds

• The experimental programme isolates the effect of the lightweight aggregate and the progressive addition of crack-bridging reinforcement.
• NWC is the structural reference. LWC00 establishes what the porous lightweight aggregate does without fibre. LWC05, LWC10 and LWC15 then trace the fibre-content response.
• LWC15 was selected for the barrier comparison because it maximized repeated-impact endurance and cost per unit of absorbed impact energy while maintaining strength above the M40 requirement.
• The source document contains an unresolved method inconsistency: one passage states that compensation water was added, another says it was not added directly, and the final table lists an extra-water quantity. At the user’s direction, this deck omits that quantity rather than silently choosing one version.
• The source also alternates between “wet” and “dry” density. Accordingly, the deck uses the neutral term “reported density.” Confirm the measurement basis before patent examples, specifications or publications are finalized.
"""
    else:
        mix_notes = """
TESTED EMBODIMENTS — approximately 60 seconds

• The programme compares a normal-weight structural reference, an unreinforced lightweight control and three progressively reinforced lightweight variants.
• The highest-performing tested variant was selected for the barrier comparison because it maximized repeated-impact endurance and laboratory cost per unit of absorbed energy while maintaining strength above M40.
• Exact material identities, dosages, moisture-conditioning protocol, mixing sequence and quality-control limits are intentionally excluded from this external version.
• The source alternates between “wet” and “dry” density; the deck therefore uses the neutral phrase “reported density.”
• Do not provide the internal formulation table or related speaker notes to an external audience until the patent filing is complete and IITBBS IP approval is documented.
"""
    add_notes(slide, mix_notes)


def slide_static(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Static performance clears the structural gate", "Measured performance", 8, status)
    add_column_chart(slide, 0.72, 1.44, 3.75, 3.78, MIXES, [("Reported density", DENSITY)],
                     [BLUE], 2700, "0", False, True, "kg/m³", 62, "Reported density")
    add_column_chart(slide, 4.78, 1.44, 3.75, 3.78, MIXES, [("Compression", COMP)],
                     [TEAL], 65, "0.0", False, True, "MPa", 62, "28 d compressive strength")
    add_column_chart(slide, 8.84, 1.44, 3.75, 3.78, MIXES, [("Split tension", TENSILE)],
                     [ORANGE], 9, "0.00", False, True, "MPa", 62, "28 d splitting tensile strength")
    add_pill(slide, 0.84, 5.49, 2.65, f"LWC15 density −{DENSITY_REDUCTION*100:.1f}%", LIGHT_BLUE, BLUE, 9.2)
    add_pill(slide, 4.93, 5.49, 2.65, f"LWC10 max = {max(COMP):.1f} MPa", LIGHT_GREEN, GREEN, 9.2)
    add_pill(slide, 8.99, 5.49, 2.65, f"LWC15 tension +{TENSILE_GAIN*100:.0f}%", LIGHT_ORANGE, ORANGE, 9.2)
    add_rect(slide, 0.85, 5.96, 11.47, 0.38, LIGHT_RED, LIGHT_RED, radius=True, line_width=0)
    add_text(slide, 1.03, 6.05, 11.05, 0.21,
             "Workability trade-off: slump fell from 165 mm (LWC00) to 87 mm (LWC15); fibre dispersion and compaction are scale-up control points.",
             8.8, RED, True, margin=0)
    add_source(slide, 0.74, 6.49, 11.65,
               "Inventor dataset, manuscript Fig. 8 and Fig. 9 native chart caches. Values shown as supplied; no independent raw-data audit. Three cube/cylinder specimens per mix were stated in the casting plan.")
    add_notes(slide, """
STATIC RESULTS — approximately 75 seconds

• The material clears the first gate: all mixtures are reported above the M40 target used for rigid crash barriers in the source.
• The selected LWC15 mixture has a reported density of 1,987 kg per cubic metre versus 2,466 for NWC—a 19.4% reduction. Because the manuscript alternates between wet and dry density, describe this as the reported value until the test basis is confirmed.
• Compressive strength is not the main source of the impact benefit, but it is preserved. LWC10 has the maximum measured value at 56.8 MPa; LWC15 remains at 55.4 MPa, 11.7% above NWC.
• The more important shift is tensile: LWC15 reaches 7.60 MPa versus 4.32 MPa for NWC, a 75.9% increase. This supports crack bridging and post-cracking integrity.
• The manufacturing trade-off is workability. Slump decreases to 87 mm at the highest fibre content. Mixing energy, fibre addition rate, moisture condition and compaction therefore belong in the process know-how and quality-control package.
""")


def slide_prism_method(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Repeated-impact method captures post-cracking endurance", "Impact methodology", 9, status)
    add_picture_contain(slide, MEDIA / "image5.jpeg", 0.73, 1.37, 5.02, 3.72, "drop_weight_setup_contain")
    add_rect(slide, 0.89, 4.43, 4.70, 0.48, NAVY, NAVY, radius=True, line_width=0)
    add_text(slide, 1.05, 4.54, 4.35, 0.24, "100 × 100 × 500 mm beam  •  400 mm support span", 9.1, WHITE, True,
             align=PP_ALIGN.CENTER, margin=0)
    # Technical parameter cards.
    params = [("13.5 kg", "drop mass"), ("100 mm", "drop height"), ("13.24 J", "energy / blow"), ("N₁ / N₂", "first crack / failure")]
    for i, (v, lab) in enumerate(params):
        x = 6.06 + (i % 2)*3.07; y = 1.43 + (i // 2)*1.30
        add_metric(slide, x, y, 2.78, 1.07, v, lab, [BLUE, CYAN, ORANGE, RED][i], WHITE, 20)
    add_rect(slide, 6.06, 4.18, 5.84, 1.27, WHITE, GRID, radius=True)
    add_text(slide, 6.28, 4.40, 1.86, 0.34, "SIGNAL PIPELINE", 9.0, BLUE, True, margin=0)
    pipeline = ["detrend", "4th-order zero-phase\nHP filter, 3 Hz", "integrate → v", "drift correct", "integrate → x"]
    for i, txt in enumerate(pipeline):
        x = 7.52 + i*0.89
        add_rect(slide, x, 4.31, 0.74, 0.72, LIGHT_BLUE if i % 2 == 0 else LIGHT_CYAN,
                 LIGHT_BLUE if i % 2 == 0 else LIGHT_CYAN, radius=True, line_width=0)
        add_text(slide, x + 0.03, 4.45, 0.68, 0.42, txt, 7.4, NAVY, True,
                 align=PP_ALIGN.CENTER, valign=MSO_VERTICAL_ANCHOR.MIDDLE, margin=0)
        if i < 4: add_chevron(slide, x + 0.73, 4.50, 0.16, 0.24, GRID)
    add_rect(slide, 6.06, 5.67, 5.84, 0.50, LIGHT_ORANGE, LIGHT_ORANGE, radius=True, line_width=0)
    add_text(slide, 6.27, 5.78, 5.42, 0.26,
             "Energy metric = mgh × number of blows; repeated impact emphasizes residual toughness.",
             9.2, INK, True, margin=0)
    add_source(slide, 0.73, 6.43, 11.65,
               "Experimental photo: inventor-supplied manuscript Fig. 4. Method: ACI 544.2R-89 adapted to prismatic beams; manuscript Testing Methods, Eqs. 1–2 and signal-processing description.")
    add_notes(slide, """
REPEATED-IMPACT METHOD — approximately 75 seconds

• The prism test uses a 13.5-kilogram hammer dropped 100 millimetres onto a simply supported 100-by-100-by-500-millimetre beam. Each impact supplies approximately 13.24 joules.
• N1 is the number of blows to the first visible crack; N2 is the number to complete failure. Multiplying the counts by mgh gives first-crack and failure energy.
• This test is especially useful for the invention because the differentiating mechanism is expected after cracking. The gap N2 minus N1 is a direct, if simplified, indicator of residual post-cracking endurance.
• An accelerometer recorded the first-impact response. The manuscript describes linear detrending, a fourth-order zero-phase Butterworth high-pass filter at 3 hertz, trapezoidal integration, drift correction and a second integration to derive displacement.
• This is technically stronger than relying only on compressive strength. However, repeated low-velocity drop-weight tests are screening tests—not a substitute for a full-scale oblique vehicle–barrier crash under MASH, EN 1317 or the applicable IRC protocol.
""")


def slide_prism_results(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide,
               "Steel fibres transform post-cracking impact endurance" if internal else
               "Crack-bridging reinforcement transforms post-cracking endurance",
               "Prism impact results", 10, status)
    add_bar_chart(slide, 0.73, 1.40, 6.02, 4.55, MIXES, PRISM_E2, BLUE, "0", 3300,
                  "Failure energy, E₂ (N·m)", [GRID, "9ABFD1", CYAN, ORANGE, RED])
    add_column_chart(slide, 7.05, 1.40, 5.48, 4.55, MIXES,
                     [("First crack N₁", PRISM_N1), ("Failure N₂", PRISM_N2)],
                     [CYAN, NAVY], 250, "0", True, True, "blows", 40, "Impact counts")
    add_pill(slide, 0.92, 5.65, 2.45, f"{PRISM_GAIN_NWC:.1f}× vs NWC", LIGHT_RED, RED, 9.3)
    add_pill(slide, 3.60, 5.65, 2.78, f"{PRISM_GAIN_LWC0:.1f}× vs plain LWC", LIGHT_ORANGE, ORANGE, 9.3)
    add_pill(slide, 7.33, 5.65, 2.22, "N₂ = 221 blows", LIGHT_BLUE, BLUE, 9.3)
    add_pill(slide, 9.78, 5.65, 2.30, "186 post-crack blows", LIGHT_CYAN, TEAL, 9.3)
    add_source(slide, 0.74, 6.42, 11.65,
               "Inventor dataset, manuscript Table 8. E₂ = mgh × N₂. Ratios calculated from supplied values; impact-count variability/replicate distribution is not reported in the table.")
    add_notes(slide, """
PRISM IMPACT RESULTS — approximately 80 seconds

• Plain lightweight concrete is more brittle than NWC in this test: LWC00 fails after two blows and absorbs 26.49 newton-metres.
• Progressive crack-bridging reinforcement changes the response nonlinearly. The three reinforced variants reach 264.9, 596.0 and 2,927.8 newton-metres at failure, respectively.
• Relative to NWC, LWC15 absorbs 55.3 times the cumulative failure energy. Relative to the plain lightweight control, the factor is 110.5. These very large ratios are driven by post-cracking endurance rather than a comparable increase in compressive strength.
• LWC15 cracks after 35 blows but does not fail until 221 blows, leaving 186 blows in the post-cracking phase. This directly supports the crack-bridging and fibre pull-out mechanism.
• Use the magnitude as strong experimental evidence, but avoid implying a universal material factor. The table does not report dispersion for N1 and N2, and drop-weight impact counts often show high variability. Repeat testing with raw specimen-level data and a suitable statistical model should be a next-step validation task.
""")


def slide_deformation(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Deformation reveals the transition from brittle to ductile response", "Dynamic response", 11, status)
    add_column_chart(slide, 0.73, 1.42, 5.38, 4.35, MIXES, [("Peak displacement", DISP)],
                     [ORANGE], 1.0, "0.00", False, True, "mm", 60, "Peak displacement after first impact")
    # Mechanism illustration. The external deck replaces constituent photos with
    # generic editable shapes so claim-enabling identity is not disclosed.
    if internal:
        add_picture_cover(slide, MEDIA / "image1.jpeg", 6.42, 1.48, 1.58, 1.26, "sfa_material", True)
        add_picture_cover(slide, MEDIA / "image3.jpeg", 8.23, 1.48, 1.58, 1.26, "steel_fibre", True)
        add_picture_contain(slide, MEDIA / "image9.jpeg", 10.04, 1.48, 2.15, 1.26, "split_failure_contain")
        labels = [(6.42, 1.58, "Porous SFA", BLUE), (8.23, 1.58, "Hooked fibres", ORANGE),
                  (10.04, 2.15, "Controlled crack path", TEAL)]
    else:
        generic = [(6.42, 1.58, LIGHT_BLUE, BLUE, "POROUS\nPHASE"),
                   (8.23, 1.58, LIGHT_ORANGE, ORANGE, "BRIDGING\nPHASE")]
        for x, w, fill, col, txt in generic:
            add_rect(slide, x, 1.48, w, 1.26, fill, fill, radius=True, line_width=0)
            add_circle_label(slide, x + w/2 - 0.24, 1.69, 0.48, "●", col, WHITE, 10)
            add_text(slide, x + 0.10, 2.25, w - 0.20, 0.34, txt, 8.4, col, True,
                     align=PP_ALIGN.CENTER, margin=0)
        add_picture_contain(slide, MEDIA / "image9.jpeg", 10.04, 1.48, 2.15, 1.26, "split_failure_external_contain")
        labels = [(6.42, 1.58, "Energy-absorbing phase", BLUE),
                  (8.23, 1.58, "Crack-bridging phase", ORANGE),
                  (10.04, 2.15, "Controlled crack path", TEAL)]
    for x, w, txt, col in labels:
        add_text(slide, x, 2.83, w, 0.36, txt, 9.3, col, True, align=PP_ALIGN.CENTER, margin=0)
    add_chevron(slide, 7.96, 1.90, 0.24, 0.34, GRID)
    add_chevron(slide, 9.77, 1.90, 0.24, 0.34, GRID)
    add_rect(slide, 6.42, 3.45, 5.77, 2.06, NAVY, NAVY, radius=True, line_width=0)
    mech = [
        ("Crush", "porous aggregate consumes local energy"),
        ("Bridge", "fibres transfer stress across opening cracks"),
        ("Pull out", "friction extends the failure process"),
        ("Hold", "reinforcement restrains large fragments"),
    ]
    for i, (a, b) in enumerate(mech):
        x = 6.72 + (i % 2)*2.72; y = 3.76 + (i // 2)*0.77
        add_text(slide, x, y, 0.80, 0.30, a.upper(), 8.6, CYAN, True, margin=0)
        add_text(slide, x + 0.78, y - 0.01, 1.70, 0.46, b, 8.6, WHITE, False, margin=0)
    add_rect(slide, 0.89, 5.87, 11.30, 0.38, LIGHT_RED, LIGHT_RED, radius=True, line_width=0)
    add_text(slide, 1.06, 5.96, 10.96, 0.21,
             "LWC10 has the highest first-impact displacement (0.85 mm); LWC15 is slightly stiffer (0.71 mm) yet vastly more durable under repeated impact.",
             8.9, RED, True, margin=0)
    add_source(slide, 0.73, 6.43, 11.65,
               "Inventor dataset, manuscript Table 9 and material/failure photos (Figs. 2 and 10). Mechanism supported by Wang & Wang (2013), Gao et al. (1997) and Sahoo et al. (2020). Error values supplied: SD 0.005–0.0686 mm.")
    mechanism_note = ("The material mechanism is a coupled sequence: SFA crushing, fibre bridging, debonding and frictional pull-out, followed by restraint of large fragments."
                      if internal else
                      "The protected material mechanism couples local crushing in an energy-absorbing phase with bridging, debonding and frictional pull-out in a reinforcing phase. Constituent identity and proportions remain redacted.")
    add_notes(slide, f"""
DEFORMATION AND MECHANISM — approximately 75 seconds

• Peak displacement after the first impact rises sharply with reinforcement. The intermediate high-reinforcement variant reaches 0.85 millimetres and the selected variant reaches 0.71, compared with 0.28 for NWC and only 0.08 for plain LWC.
• The lower displacement of plain LWC should not be interpreted as superior stiffness. In the context of its immediate failure, it indicates a brittle, low-deformation response.
• The selected variant is slightly less deformable than the preceding variant on the first impact, plausibly because the denser reinforcement network increases stiffness. Yet it has vastly greater repeated-impact endurance. First-impact displacement and cumulative post-cracking toughness are different response dimensions.
• {mechanism_note}
• The source supplies standard deviations for displacement; a future analysis should present full time histories, filtering sensitivity, uncertainty bands and independent displacement verification, because double integration of acceleration is sensitive to drift and cutoff selection.
""")


def slide_economics(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "The optimum is expensive per m³—but efficient per absorbed joule", "Techno-economics", 12, status)
    add_column_chart(slide, 0.74, 1.47, 5.62, 4.25, MIXES, [("Material cost", COST)],
                     [NAVY], 18000, "₹#,##0", False, True, "₹/m³", 58, "Material cost")
    add_column_chart(slide, 6.76, 1.47, 5.62, 4.25, MIXES, [("Cost / failure energy", COST_E)],
                     [ORANGE], 185, "0.00", False, True, "₹/(N·m)", 58, "Laboratory cost / impact energy")
    add_pill(slide, 0.94, 5.72, 2.50, f"LWC15 cost = {COST[-1]/COST[0]:.2f}× NWC", LIGHT_RED, RED, 9.1)
    add_pill(slide, 3.75, 5.72, 2.25,
             "Steel fibre dominates cost" if internal else "Reinforcement dominates cost",
             LIGHT_ORANGE, ORANGE, 9.1)
    add_pill(slide, 7.02, 5.72, 2.68, f"Unit impact value = {COST_ENERGY_GAIN:.1f}× better", LIGHT_CYAN, TEAL, 9.1)
    add_pill(slide, 10.00, 5.72, 2.10, "₹5.53 per N·m", LIGHT_BLUE, BLUE, 9.1)
    add_source(slide, 0.74, 6.43, 11.65,
               "Inventor dataset, manuscript Tables 7 and 10. Local material prices × mix quantities; cost/energy uses prism E₂. Excludes labour, transport, curing, reinforcement, fabrication, repair, service life, inflation and test uncertainty.")
    cost_driver = ("the high steel-fibre dosage" if internal else "the high reinforcement dosage")
    add_notes(slide, f"""
TECHNO-ECONOMIC RESULT — approximately 80 seconds

• The selected lightweight variant is not the cheapest concrete. Its inventor-estimated material cost is about 16,192 rupees per cubic metre, 3.27 times NWC, primarily because of {cost_driver}.
• The relevant engineering denominator is failure energy. When the reported prism impact capacity is included, LWC15 costs 5.53 rupees per newton-metre versus 93.36 for NWC—about 16.9 times better on this laboratory metric.
• This supports selection of LWC15 for the barrier test and creates a compelling value narrative: pay more per unit volume, receive substantially more post-cracking impact capacity.
• Keep the boundary explicit. This is not a full barrier cost, lifecycle cost or cost-benefit analysis. It omits reinforcement, formwork, labour, transport, curing, installation, repair frequency, downtime and uncertainty. Prices are local and time-specific.
• The commercial study should therefore compare complete installed systems and include lower transport mass, bridge-deck dead load, expected repair after impact, quality-control cost, industrial-aggregate logistics and any savings from modular precast deployment.
""")


def slide_barrier_method(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Application-scale test: 1:3 New Jersey barrier under pendulum impact", "Scaled validation", 13, status)
    add_picture_contain(slide, MEDIA / "image6.png", 0.72, 1.34, 5.65, 3.82, "barrier_geometry_contain")
    add_picture_contain(slide, MEDIA / "image8.jpeg", 6.67, 1.34, 5.65, 3.82, "pendulum_setup_contain")
    add_text(slide, 0.92, 4.77, 5.25, 0.29, "Full-scale and 1:3 cross-section / reinforcement geometry", 8.5, INK, True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, 6.87, 4.77, 5.25, 0.29, "Pendulum rig: restrained specimen; rear-face accelerometer", 8.5, INK, True,
             align=PP_ALIGN.CENTER, margin=0)
    vals = [("40 kg", "impactor"), ("0.35 m", "vertical rise"), ("≈137.3 J", "energy / blow"), ("1 each", "barrier / mix")]
    for i, (v, lab) in enumerate(vals):
        x = 0.84 + i*2.93
        add_metric(slide, x, 5.30, 2.60, 0.88, v, lab, [BLUE, CYAN, ORANGE, RED][i], WHITE, 17)
    add_source(slide, 0.73, 6.42, 11.65,
               "Inventor-supplied manuscript Figs. 5–7; Testing Methods. Geometry scaled 1:3 using dimensional similarity; energy per blow = 40 × 9.81 × 0.35 ≈ 137.3 J. One NWC and one LWC15 barrier were tested.")
    add_notes(slide, """
SCALED BARRIER TEST — approximately 80 seconds

• To move beyond material coupons, the team fabricated one-third-scale New Jersey-profile barriers using NWC and the selected LWC15 mixture.
• The dimensions were reduced by a factor of three, and a pendulum rig applied repeated impact at mid-height. The impactor mass was 40 kilograms, the vertical rise 0.35 metres and the energy approximately 137.3 joules per blow.
• An accelerometer on the rear face captured the response for the same filtering and integration workflow used in the beam tests.
• This experiment is valuable as an application-scale proof of concept and supports reduction to practice. It also provides comparative failure patterns and cumulative impact energy.
• It is not a standards-compliant crash test. Geometric similarity alone does not guarantee similitude of strain rate, fracture, reinforcement bond, contact, mass, velocity and gravity effects. The specimen was restrained, and there was no vehicle to evaluate containment, redirection, wheel climb, rollover, occupant compartment intrusion, ASI or THIV.
• Also note the sample size: one barrier per material. Treat the result as preliminary but decision-relevant.
""")


def slide_barrier_results(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Scaled-barrier validation preserves the direction of benefit", "Barrier results", 14, status)
    add_picture_contain(slide, MEDIA / "image16.jpeg", 0.72, 1.37, 5.70, 4.53, "barrier_failure_contain")
    add_column_chart(slide, 6.75, 1.40, 2.77, 3.15, BARRIER_MIXES, [("Failure energy", BARRIER_E2)],
                     [RED], 3100, "0", False, True, "N·m", 55, "Energy to failure")
    add_column_chart(slide, 9.78, 1.40, 2.77, 3.15, BARRIER_MIXES, [("Deformation", BARRIER_DEF)],
                     [ORANGE], 15, "0.00", False, True, "mm", 55, "Peak deformation")
    add_metric(slide, 6.75, 4.82, 1.78, 1.12, f"{BARRIER_ENERGY_GAIN:.1f}×", "failure energy", RED, WHITE, 20)
    add_metric(slide, 8.66, 4.82, 1.78, 1.12, f"{BARRIER_DEF_GAIN:.2f}×", "deformation", ORANGE, WHITE, 20)
    add_metric(slide, 10.57, 4.82, 1.98, 1.12, "20 vs 10", "blows to failure", TEAL, WHITE, 20)
    add_rect(slide, 0.82, 6.02, 11.65, 0.35, LIGHT_RED, LIGHT_RED, radius=True, line_width=0)
    add_text(slide, 1.00, 6.10, 11.25, 0.20,
             "Fragmentation observation is qualitative: LWC15 showed finer, more distributed debris; particle-size/mass and ejection-velocity measurements are still required.",
             8.6, RED, True, margin=0)
    add_source(slide, 0.73, 6.45, 11.65,
               "Inventor dataset, manuscript Tables 11–12 and Fig. 16. NWC: N₂=10, E₂=1,373 N·m, 4.22 mm; LWC15: N₂=20, E₂=2,746 N·m, 12.70 mm. n=1 barrier per mix.")
    add_notes(slide, """
BARRIER RESULTS — approximately 90 seconds

• The application-scale result points in the same direction as the prism tests. The NWC barrier failed after 10 blows and 1,373 newton-metres; LWC15 reached 20 blows and 2,746 newton-metres—exactly twice the cumulative impact energy in the reported setup.
• Peak deformation increased from 4.22 to 12.70 millimetres, a factor of approximately 3.01. Controlled deformation is consistent with the intended energy-dissipation philosophy.
• The failure photographs show a broader, more brittle crack and large localized fragments in NWC, while LWC15 exhibits a narrower crack path and finer, more distributed debris after more impacts.
• That fragmentation finding is qualitative. A stronger patent and commercialization evidence package should sieve and weigh debris, map crack width and density, measure ejection velocity, and assess residual load capacity after prescribed impacts.
• Because n equals one per mix, do not attach confidence intervals or claim a guaranteed two-times product performance. State that the single comparative barrier experiment reproduced the direction of the material-level benefit and justifies replicated/full-scale validation.
""")


def slide_evidence_claims(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Measured evidence maps to a layered patent strategy", "IP evidence map", 15, status)
    rows = [
        ["Evidence", "Observed LWC15 result", "Supports", "Confidence"],
        ["Reported density", f"{DENSITY[-1]:,.0f} kg/m³; −{DENSITY_REDUCTION*100:.1f}% vs NWC", "Lightweight article / transport value", "Medium*"],
        ["Compressive strength", f"{COMP[-1]:.1f} MPa; M40+", "Structural-use enablement", "High"],
        ["Split tensile", f"{TENSILE[-1]:.2f} MPa; +{TENSILE_GAIN*100:.0f}%", "Fibre-bridging function", "High"],
        ["Prism failure energy", f"{PRISM_E2[-1]:,.1f} N·m; {PRISM_GAIN_NWC:.1f}× NWC", "Unexpected technical effect", "Medium"],
        ["Scaled barrier", f"{BARRIER_E2[-1]:,.0f} N·m; {BARRIER_ENERGY_GAIN:.1f}× NWC", "Barrier-use embodiment", "Preliminary"],
        ["Fragmentation", "Finer / distributed debris observed", "Controlled-damage function", "Preliminary"],
    ]
    add_table(slide, 0.72, 1.32, 7.45, 4.65, rows, [1.55, 2.35, 2.15, 1.05], font_size=8.6,
              first_col_left=True, highlight_rows={4: LIGHT_ORANGE, 5: LIGHT_CYAN})
    # Claim stack.
    add_rect(slide, 8.47, 1.33, 3.92, 4.65, NAVY, NAVY, radius=True, line_width=0)
    add_text(slide, 8.78, 1.63, 3.28, 0.40,
             "ILLUSTRATIVE CLAIM STACK" if internal else "PROTECTABLE VALUE STACK",
             10.2, CYAN, True, margin=0)
    if internal:
        claims = [
            ("A", "Composition", "OPC / GGBS / SFA / hooked-fibre ranges"),
            ("B", "Process", "conditioning, sequence, dispersion and QC"),
            ("C", "Barrier article", "profile + reinforcement + invented concrete"),
            ("D", "Performance", "density / strength / impact / damage envelope"),
            ("E", "System use", "precast transport-safety applications"),
        ]
    else:
        claims = [
            ("A", "Material platform", "integrated lightweight energy-dissipation system"),
            ("B", "Manufacturing know-how", "controlled processing and quality assurance"),
            ("C", "Barrier embodiment", "application-specific structural article"),
            ("D", "Performance envelope", "strength, impact endurance and damage control"),
            ("E", "Deployment package", "precast production and field implementation"),
        ]
    for i, (letter, title, body) in enumerate(claims):
        y = 2.19 + i*0.67
        add_circle_label(slide, 8.80, y, 0.38, letter, [BLUE, CYAN, TEAL, ORANGE, RED][i], WHITE, 9)
        add_text(slide, 9.30, y - 0.01, 1.12, 0.27, title, 9.4, WHITE, True, margin=0)
        add_text(slide, 10.43, y - 0.03, 1.59, 0.42, body, 7.9, "C4D4DC", margin=0)
    add_text(slide, 8.78, 5.56, 3.25, 0.30, "* Density basis must be corrected.", 7.8, ORANGE, True, margin=0)
    add_source(slide, 0.73, 6.35, 11.65,
               "Evidence: inventor dataset. Claim structure is a technical drafting aid only—not legal advice. Patentability/FTO requires professional searching, claim construction and jurisdiction-specific counsel.")
    if internal:
        claim_notes = """
EVIDENCE-TO-CLAIM MAP — approximately 90 seconds

• This slide separates what is measured from what it may support in a patent application.
• Strength and splitting tensile data provide the strongest conventional evidence. The repeated-impact result may support an unexpected technical effect, but its uncertainty and test adaptation should be disclosed accurately.
• The barrier-scale result supports an article-of-manufacture embodiment and utility, while remaining preliminary due to n equals one and the absence of a standards-compliant vehicle test.
• A layered filing strategy can include: composition ranges; a process claim around moisture conditioning, mixing sequence, fibre dispersion and quality control; a barrier article made from the material; and carefully drafted functional/performance limitations where enabled.
• Avoid relying only on a narrow point recipe; include supported ranges and fallback positions around binder fraction, SFA substitution, fibre geometry/dosage and performance windows. Do not invent ranges not supported by the notebooks and experiments.
• Before filing, reconcile the water protocol and density terminology, preserve raw data and notebooks, and complete a professional patentability/FTO search. This slide is a technical aid and not legal advice.
"""
    else:
        claim_notes = """
EVIDENCE-TO-VALUE MAP — approximately 75 seconds

• This slide separates measured performance from the protectable value categories that may be relevant after filing.
• Static strength, tensile resistance, repeated-impact energy, scaled-barrier response and observed damage control together support a coherent technology platform.
• The barrier result remains preliminary because only one barrier per material was tested and no standards-compliant vehicle test was performed.
• Detailed composition ranges, processing parameters, conditioning controls and claim fallbacks are intentionally excluded from this external presentation.
• Do not answer external questions about formulation or manufacturing specifics from memory. Refer them to the IITBBS technology-transfer/IP team after filing and approval.
• The claim/value structure shown here is a technical communication aid, not legal advice or a patentability assurance.
"""
    add_notes(slide, claim_notes)


def slide_translation(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "De-risking plan: convert a laboratory signal into a certifiable product", "Translation roadmap", 16, status)
    phases = [
        ("NOW", "Protect", "File before disclosure; professional search; lock data provenance", BLUE),
        ("0–6 m", "Reproduce", "Resolve method conflicts; replicate barrier tests; quantify debris", CYAN),
        ("6–12 m", "Model + age", "Calibrate FE; fatigue/repeat impact; chloride, fire and weather", TEAL),
        ("12–24 m", "Full-scale certify", "Vehicle crash: containment, redirection, ASI/THIV and debris", ORANGE),
        ("24–36 m", "Pilot + license", "Precast QA, field demonstration, LCA/TEA and partner transfer", RED),
    ]
    for i, (time, title, body, col) in enumerate(phases):
        x = 0.73 + i*2.46
        add_text(slide, x, 1.36, 2.12, 0.28, time, 8.8, col, True, margin=0,
                 align=PP_ALIGN.CENTER)
        add_circle_label(slide, x + 0.80, 1.75, 0.52, str(i+1), col, WHITE, 10)
        if i < 4:
            add_line(slide, x + 1.32, 2.01, x + 2.47, 2.01, GRID, 2.1)
        add_text(slide, x, 2.51, 2.12, 0.34, title, 12.5, INK, True, FONT_HEAD,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, x + 0.06, 2.96, 2.00, 1.12, body, 9.0, MUTED, False,
                 align=PP_ALIGN.CENTER, margin=0)
    # Risk bar.
    add_rect(slide, 0.75, 4.52, 11.82, 1.48, NAVY, NAVY, radius=True, line_width=0)
    add_text(slide, 1.02, 4.80, 1.25, 0.30, "TOP RISKS", 9.2, CYAN, True, margin=0)
    risks = [
        "Full-scale vehicle response", "Fibre workability / balling" if internal else "Reinforcement workability",
        "SFA moisture + variability" if internal else "Aggregate moisture + variability",
        "Durability / corrosion", "Installed-system cost", "Patent landscape / disclosure timing"
    ]
    for i, r in enumerate(risks):
        x = 2.26 + (i % 3)*3.29; y = 4.72 + (i // 3)*0.58
        add_circle_label(slide, x, y + 0.01, 0.28, "!", RED, WHITE, 8)
        add_text(slide, x + 0.39, y, 2.70, 0.34, r, 8.7, WHITE, True, margin=0)
    add_source(slide, 0.73, 6.34, 11.65,
               "Roadmap synthesizes inventor questionnaire/technology profile and validation gaps against IRC:119-2015, EN 1317/MASH-style evaluation principles, and peer-reviewed barrier crash-test literature. Indicative timing only.")
    pilot_controls = ("precast quality controls, fibre-dispersion inspection, SFA moisture/specification limits"
                      if internal else
                      "precast quality controls and protected material/process acceptance limits")
    add_notes(slide, f"""
TRANSLATION ROADMAP — approximately 90 seconds

• The immediate action is protection: reconcile the technical record, complete professional searching and file before any publication, conference presentation, industry pitch or external distribution.
• In the first six months, repeat the key tests with specimen-level data, quantify uncertainty, resolve the method and density conflicts, replicate barrier specimens and turn fragmentation into a measured particle-size, mass and velocity dataset.
• In parallel, calibrate an LS-DYNA or ABAQUS vehicle–barrier model against the experiments, and add durability, corrosion, environmental exposure, repeated-impact and repairability studies.
• The decisive commercialization gate is a full-scale vehicle crash conducted to the required jurisdictional standard. It must evaluate containment, redirection, stability, debris, occupant compartment performance and occupant-risk measures such as ASI or THIV where applicable.
• A pilot then needs {pilot_controls}, a full installed-system cost model and comparative LCA.
• Candidate partners listed by the inventors include CRRI, MoRTH/NHAI, state PWDs, precast manufacturers, barrier suppliers and major infrastructure contractors.
""")


def slide_decision(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Decision request: approve patent filing and staged validation", "Recommendation", 17, status, dark=True)
    add_text(slide, 0.75, 1.36, 7.65, 0.95,
             "The evidence justifies IPR approval to proceed with a focused patent filing—\nnot yet a certified occupant-safety claim.",
             23.0, WHITE, True, FONT_HEAD, margin=0, line_spacing=0.93)
    # Big decision metrics.
    vals = [(f"{PRISM_GAIN_NWC:.1f}×", "prism E₂ vs NWC", ORANGE),
            (f"{BARRIER_ENERGY_GAIN:.1f}×", "scaled barrier E₂", RED),
            (f"{COST_ENERGY_GAIN:.1f}×", "lab unit-impact value", CYAN)]
    for i, (v, lab, col) in enumerate(vals):
        x = 0.80 + i*2.55
        add_rect(slide, x, 2.65, 2.28, 1.32, "0E3047", "20485F", radius=True, line_width=0.8)
        add_text(slide, x + 0.17, 2.85, 1.94, 0.44, v, 24, col, True, FONT_HEAD, margin=0,
                 align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.17, 3.40, 1.94, 0.28, lab, 8.9, "C3D5DD", True, margin=0,
                 align=PP_ALIGN.CENTER)
    add_rect(slide, 8.78, 1.36, 3.66, 4.87, WHITE, WHITE, radius=True, line_width=0)
    add_text(slide, 9.10, 1.72, 3.00, 0.40, "IPR DECISIONS REQUESTED", 10.2, BLUE, True, margin=0)
    asks = [
        ("1", "Approve patent filing", "Authorize search, drafting and filing"),
        ("2", "Maintain confidentiality", "Coordinate manuscript release after filing"),
        ("3", "Endorse validation plan", "Replicates, debris metrics, FE and full-scale crash"),
        ("4", "Enable technology transfer", "Controlled CRRI / NHAI / precast engagement"),
    ]
    for i, (n, t, b) in enumerate(asks):
        y = 2.31 + i*0.91
        add_circle_label(slide, 9.10, y, 0.42, n, [BLUE, ORANGE, TEAL, RED][i], WHITE, 9)
        add_text(slide, 9.66, y - 0.02, 2.48, 0.36, t, 9.4, INK, True, margin=0)
        add_text(slide, 9.66, y + 0.34, 2.38, 0.38, b, 7.9, MUTED, margin=0)
    add_rect(slide, 0.80, 4.42, 7.48, 1.28, "0E3047", "20485F", radius=True, line_width=0.8)
    add_text(slide, 1.06, 4.70, 1.20, 0.30, "BOTTOM LINE", 9.0, CYAN, True, margin=0)
    add_text(slide, 2.22, 4.61, 5.75, 0.64,
             "Approve filing now; retain full-scale occupant-risk validation as the commercial release gate.",
             12.2, WHITE, True, margin=0)
    add_source(slide, 0.80, 6.44, 11.60,
               "Recommendation based on inventor-supplied evidence and inventor-confirmed pre-filing status as of 13 Aug 2026. No legal, safety-certification or investment assurance is implied.", True)
    add_notes(slide, """
CLOSE — approximately 60 seconds

• The recommendation to the IIT Bhubaneswar IPR section is to approve proceeding with a focused patent filing. The team has reduced the concept to practice and demonstrated a distinctive combination of lower reported density, retained structural strength, major repeated-impact endurance and a positive one-third-scale barrier result.
• The evidence does not yet support a certified occupant-safety claim. The commercial release gate must be a full-scale standards-compliant vehicle test with occupant-risk, containment, redirection, stability and debris criteria.
• Four IPR decisions are requested: approve professional prior-art searching, drafting and filing; maintain confidentiality and coordinate manuscript release; endorse a staged validation programme; and enable controlled engagement with crash-test and precast partners.
• The claim strategy should protect the integrated platform rather than a single recipe: material composition and ranges, process controls, a barrier article, and the use/performance relationship where legally enabled.
• Final line: approve protection now, validate the occupant benefit next, and commercialize only after the full-scale safety gate is passed.
""")


def slide_references_1(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "References & primary sources (1/2)", "Sources", 18, status)
    primary_article_ref = ("1  Al-Amir, H.; Kanungo, P.; Chandrappa, A.K.; Pasla, D. (2026). Novel Steel-Fiber Reinforced Sintered Fly Ash Light Weight Aggregate Concrete for Crash Barrier Applications. Inventor-supplied manuscript; submitted 18 Jul 2026."
                           if internal else
                           "1  IIT Bhubaneswar invention team (2026). Confidential inventor-supplied crash-barrier research manuscript; submitted 18 Jul 2026. Title and formulation details redacted pending filing.")
    refs_left = [
        primary_article_ref,
        "2  IIT Bhubaneswar invention team (2026). Invention and Technology Disclosure Form: Novel Concrete Composition for Impact Energy Dissipating Crash Barriers, dated 16 Jul 2026; Questionnaire; Technology Profile.",
        "3  World Health Organization (2023). Global Status Report on Road Safety 2023. https://www.who.int/publications/i/item/9789240086517",
        "4  UNEP (2022). Sand and Sustainability: 10 Strategic Recommendations to Avert a Crisis. https://wedocs.unep.org/20.500.11822/38362",
        "5  U.S. Geological Survey (2025). Mineral Commodity Summaries 2025—Cement. https://pubs.usgs.gov/periodicals/mcs2025/mcs2025-cement.pdf",
        "6  International Energy Agency (2023). Cement—Tracking Clean Energy Progress. https://www.iea.org/energy-system/industry/cement",
    ]
    refs_right = [
        "7  Reid, I.; Carpenter, A.M.; Masili, A. (2020). Beneficial Uses of Coal Fly Ash. IEA Clean Coal Centre, CCC/303, ISBN 978-92-9029-626-3.",
        "8  IRC:119-2015. Guidelines for Traffic Safety Barriers. Indian Roads Congress, New Delhi.",
        "9  IRC:6-2016. Standard Specifications and Code of Practice for Road Bridges. Indian Roads Congress.",
        "10 ACI 211.2-98. Standard Practice for Selecting Proportions for Structural Lightweight Concrete. American Concrete Institute.",
        "11 ACI 544.2R-89. Measurement of Properties of Fiber Reinforced Concrete. American Concrete Institute.",
        ("12 IS 9142 (Part 2):2018. Artificial Lightweight Aggregates—Sintered Fly Ash Aggregate. Bureau of Indian Standards."
         if internal else
         "12 Applicable Bureau of Indian Standards lightweight-aggregate specification; detailed material designation withheld pending patent filing."),
        "13 IIT Bhubaneswar visual identity/logo: https://www.iitbbs.ac.in/ (accessed Aug 2026).",
    ]
    for col, refs in enumerate([refs_left, refs_right]):
        x = 0.73 if col == 0 else 6.75
        add_rect(slide, x, 1.30, 5.82, 5.62, WHITE, GRID, radius=True)
        y = 1.55
        for ref in refs:
            h = 0.76 if len(ref) < 180 else 0.92
            add_text(slide, x + 0.22, y, 5.38, h, ref, 8.9, INK, False, margin=0, line_spacing=1.02)
            y += h + 0.08
    add_notes(slide, """
REFERENCE SLIDE — do not normally read aloud.

• These are the primary inventor documents, global datasets and design/test standards used in the presentation.
• The inventor manuscript is not represented as published or peer reviewed. It is identified as an inventor-supplied manuscript submitted on 18 July 2026.
• The WHO global road-death distribution, UNEP materials estimate, USGS cement-production estimate and IEA cement-emissions context are independently sourced.
• Confirm the current editions and legal applicability of IRC, BIS, ACI, MASH and EN 1317 requirements with the intended certification authority before test planning or product claims.
""")


def slide_references_2(prs, internal):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); status = status_text(internal)
    add_header(slide, "Peer-reviewed literature & proceedings (2/2)", "Sources", 19, status)
    refs_left = [
        "14 Raj, A.; Nagarajan, P.; Pallikkara, S.A. (2020). Application of Fiber-Reinforced Rubcrete for Crash Barriers. J. Mater. Civ. Eng. 32(12), 04020358. doi:10.1061/(ASCE)MT.1943-5533.0003454.",
        "15 Shen, W. et al. (2014). Investigation on the safety concrete for highway crash barrier. Constr. Build. Mater. 70, 394–398. doi:10.1016/j.conbuildmat.2014.07.060.",
        "16 Zain, M.F.M.; Mohammed, H.J. (2015). Concrete road barriers subjected to impact loads: an overview. Lat. Am. J. Solids Struct. 12, 1824–1858. doi:10.1590/1679-78251783.",
        "17 Wang, H.T.; Wang, L.C. (2013). Static and dynamic properties of steel fiber reinforced lightweight aggregate concrete. Constr. Build. Mater. 38, 1146–1151. doi:10.1016/j.conbuildmat.2012.09.016.",
        "18 Sahoo, S.; Selvaraju, A.K.; Suriya Prakash, S. (2020). Mechanical characterization of structural lightweight aggregate concrete made with SFA and synthetic fibres. Cem. Concr. Compos. 113, 103712. doi:10.1016/j.cemconcomp.2020.103712.",
        "19 Nadesan, M.S.; Dinakar, P. (2017). Mix design and properties of fly ash waste lightweight aggregates in structural lightweight concrete. Case Stud. Constr. Mater. 7, 336–347. doi:10.1016/j.cscm.2017.09.005.",
        "20 Kim, W. et al. (2019). Evaluation of concrete barriers with novel shock absorbers subjected to impact loading. Arch. Civ. Mech. Eng. 19, 657–671. doi:10.1016/j.acme.2019.01.004.",
    ]
    refs_right = [
        "21 Yang, J. et al. (2019). Crash performance evaluation of a new movable median guardrail. Eng. Struct. 182, 459–472. doi:10.1016/j.engstruct.2018.12.090.",
        "22 Abid, S.R. et al. (2020). Repeated drop-weight impact tests on self-compacting concrete reinforced with microsteel fiber. Heliyon 6, e03198. doi:10.1016/j.heliyon.2020.e03198.",
        "23 Li, N.; Park, B.B.; Lambert, J.H. (2018). Effect of guardrail on reducing fatal and severe injuries on freeways. J. Transp. Saf. Secur. 10(5), 455–470. doi:10.1080/19439962.2017.1297970.",
        "24 Grzebieta, R.H. et al. (2005). Roadside Hazard and Barrier Crashworthiness Issues. Proc. 19th Int’l Technical Conf. on Enhanced Safety of Vehicles, Paper 05-0149-O, NHTSA.",
        "25 Bonin, G.; Cantisani, G.; Loprencipe, G.; Ranzo, A. (2004). Road safety barriers with short elements of lightweight concrete. Proc. II Int’l Congress SIIV—New Technologies and Modelling Tools for Roads.",
        "26 Iqbal, M. (2011). Vehicular Impact Loading and Barrier Design. ACI Symposium Publication SP-281, 1–16. doi:10.14359/51683612.",
        "27 Pachocki, Ł.; Bruski, D. (2020). Modeling, simulation and validation of a TB41 concrete restraint-system crash test. Arch. Civ. Mech. Eng. doi:10.1007/s43452-020-00065-7.",
    ]
    for col, refs in enumerate([refs_left, refs_right]):
        x = 0.73 if col == 0 else 6.75
        add_rect(slide, x, 1.30, 5.82, 5.62, WHITE, GRID, radius=True)
        y = 1.52
        for ref in refs:
            h = 0.72 if len(ref) < 200 else 0.82
            add_text(slide, x + 0.22, y, 5.38, h, ref, 8.1, INK, False, margin=0, line_spacing=0.98)
            y += h + 0.055
    add_notes(slide, """
REFERENCE SLIDE — do not normally read aloud.

• This list deliberately includes both reputable journal literature and conference/symposium proceedings relevant to concrete-barrier crashworthiness and lightweight material concepts.
• Bonin et al. is particularly important for novelty framing because it shows that lightweight-concrete road-safety barrier concepts existed before this invention. The proposed patent position must therefore be narrower and evidence-led.
• The literature supports the plausibility of fibre bridging, lightweight-aggregate energy dissipation and the importance of occupant-risk metrics; it does not independently validate the exact invented formulation.
• A patent professional should expand this literature review into a jurisdiction-specific patentability and freedom-to-operate search before claims are finalized.
""")


def build_deck(internal: bool, output: Path) -> None:
    mode = "internal confidential" if internal else "external redacted"
    prs = base_prs(
        "Novel Concrete Composition for Impact Energy Dissipating Crash Barriers",
        f"Editable {mode} patent-value presentation with experimental evidence and references",
        "Generated from inventor-supplied documents. Claims and certification statements are evidence-qualified."
    )
    title_slide(prs, internal)
    slide_executive(prs, internal)
    slide_global(prs, internal)
    slide_prior_art(prs, internal)
    slide_invention(prs, internal)
    slide_program(prs, internal)
    slide_mix(prs, internal)
    slide_static(prs, internal)
    slide_prism_method(prs, internal)
    slide_prism_results(prs, internal)
    slide_deformation(prs, internal)
    slide_economics(prs, internal)
    slide_barrier_method(prs, internal)
    slide_barrier_results(prs, internal)
    slide_evidence_claims(prs, internal)
    slide_translation(prs, internal)
    slide_decision(prs, internal)
    slide_references_1(prs, internal)
    slide_references_2(prs, internal)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def write_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active; ws.title = "README"
    ws.append(["Patent-project presentation data workbook"])
    ws.append(["Source", "Inventor-supplied manuscript/disclosure unless otherwise stated"])
    ws.append(["Caveat", "Values transcribed from supplied documents; no independent raw-data audit. Extra-water protocol omitted pending correction; density basis unresolved."])
    ws.append(["Patent status", "Not yet filed; presentation prepared for IIT Bhubaneswar IPR committee pre-filing evaluation and approval."])
    ws.append(["Generated", "13 August 2026"])

    ws = wb.create_sheet("Global context")
    ws.append(["Metric", "Value", "Unit", "Year", "Source"])
    rows = [
        ["Global road deaths", 1.19, "million/year", 2021, "WHO Global Status Report on Road Safety 2023"],
        ["Road deaths in upper-middle income", 35, "%", 2021, "WHO 2023"],
        ["Road deaths in lower-middle income", 44, "%", 2021, "WHO 2023"],
        ["Road deaths in low income", 13, "%", 2021, "WHO 2023"],
        ["Road deaths in high income", 8, "%", 2021, "WHO 2023"],
        ["Sand and gravel use", 50, "Gt/year", 2022, "UNEP Sand and Sustainability 2022"],
        ["World cement production", 4.0, "Gt/year", 2024, "USGS MCS 2025"],
        ["Direct cement emissions intensity", 0.6, "tCO2/t cement (just under)", 2022, "IEA Cement 2023"],
        ["Global coal fly ash production", 1.0, ">Gt/year", 2020, "Reid, Carpenter & Masili, IEA Clean Coal Centre CCC/303"],
    ]
    for r in rows: ws.append(r)

    ws = wb.create_sheet("Mixture performance")
    ws.append(["Mix", "Fibre vol %", "Reported density kg/m3", "Slump mm", "Compressive MPa", "Split tensile MPa", "N1 blows", "N2 blows", "E1 N-m", "E2 N-m", "Peak disp mm", "SD mm", "Cost Rs/m3", "Cost per energy Rs/N-m"])
    for i, mix in enumerate(MIXES):
        ws.append([mix, FIBRE_VOL[i], DENSITY[i], SLUMP[i], COMP[i], TENSILE[i], PRISM_N1[i], PRISM_N2[i], PRISM_E1[i], PRISM_E2[i], DISP[i], DISP_SD[i], COST[i], COST_E[i]])

    ws = wb.create_sheet("Internal formulation")
    ws.append(["Mix"] + [f"{x} kg/m3" for x in COMPONENT_NAMES])
    for mix in MIXES: ws.append([mix] + MIX_COMPONENTS[mix])
    ws.append([])
    ws.append(["Note", "Extra-water quantity excluded by user instruction because narrative and final mix table conflict."])

    ws = wb.create_sheet("Barrier performance")
    ws.append(["Mix", "N1 blows", "N2 blows", "E1 N-m", "E2 N-m", "Peak deformation mm", "Specimen count"])
    for i, mix in enumerate(BARRIER_MIXES):
        ws.append([mix, BARRIER_N1[i], BARRIER_N2[i], BARRIER_E1[i], BARRIER_E2[i], BARRIER_DEF[i], 1])

    ws = wb.create_sheet("Calculated metrics")
    ws.append(["Metric", "Value", "Formula / basis"])
    calc = [
        ["LWC15 reported density reduction vs NWC", DENSITY_REDUCTION, "1 - 1987/2466"],
        ["LWC15 compressive gain vs NWC", COMP_GAIN, "55.4/49.62 - 1"],
        ["LWC15 tensile gain vs NWC", TENSILE_GAIN, "7.60/4.32 - 1"],
        ["LWC15 prism E2 factor vs NWC", PRISM_GAIN_NWC, "2927.81/52.97"],
        ["LWC15 prism E2 factor vs LWC00", PRISM_GAIN_LWC0, "2927.81/26.49"],
        ["LWC15 cost/energy improvement vs NWC", COST_ENERGY_GAIN, "93.36/5.53"],
        ["LWC15 barrier E2 factor vs NWC", BARRIER_ENERGY_GAIN, "2746/1373"],
        ["LWC15 barrier deformation factor vs NWC", BARRIER_DEF_GAIN, "12.70/4.22"],
    ]
    for r in calc: ws.append(r)
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=NAVY)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2): row[0].number_format = "0.000"

    ws = wb.create_sheet("Figure source register")
    ws.append(["Slide", "Figure / chart", "Source"])
    register = [
        [3, "Global road deaths doughnut", "WHO Global Status Report on Road Safety 2023, pp. 4 and 14"],
        [3, "Global context metrics", "UNEP 2022; USGS MCS 2025; IEA Cement 2023; Reid, Carpenter & Masili 2020"],
        [4, "Prior-art matrix", "Zain & Mohammed 2015; Raj et al. 2020; Shen et al. 2014; Bonin et al. 2004; Grzebieta et al. 2005"],
        [5, "Mechanism schematic", "Inventor disclosure; Wang & Wang 2013; Nadesan & Dinakar 2017; Sahoo et al. 2020"],
        [7, "Mix design", "Inventor manuscript Tables 4–6"],
        [8, "Density/compression/tension charts", "Inventor manuscript Fig. 8–9 embedded chart caches"],
        [9, "Drop-weight setup", "Inventor manuscript Fig. 4"],
        [10, "Prism impact charts", "Inventor manuscript Table 8"],
        [11, "Displacement chart/photos", "Inventor manuscript Table 9, Figs. 2 and 10"],
        [12, "Cost charts", "Inventor manuscript Tables 7 and 10"],
        [13, "Scaled barrier geometry/setup", "Inventor manuscript Figs. 5–7"],
        [14, "Barrier results/photo", "Inventor manuscript Tables 11–12 and Fig. 16"],
        [15, "Evidence/claim map", "Inventor dataset; author synthesis"],
    ]
    for r in register: ws.append(r)

    # Global formatting.
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(vertical="center")
        for col in ws.columns:
            letter = col[0].column_letter
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    if not all(p.exists() for p in [ARTICLE, DISCLOSURE, QUESTIONNAIRE, TECH_PROFILE, LOGO]):
        missing = [str(p) for p in [ARTICLE, DISCLOSURE, QUESTIONNAIRE, TECH_PROFILE, LOGO] if not p.exists()]
        raise FileNotFoundError("Missing required inputs: " + ", ".join(missing))
    BUILD.mkdir(exist_ok=True); extract_media(); OUT.mkdir(exist_ok=True)
    internal = OUT / "IITBBS_Impact_Dissipating_Crash_Barrier_Internal_Confidential.pptx"
    external = OUT / "IITBBS_Impact_Dissipating_Crash_Barrier_External_Redacted.pptx"
    workbook = OUT / "IITBBS_Crash_Barrier_Presentation_Data.xlsx"
    build_deck(True, internal)
    build_deck(False, external)
    write_workbook(workbook)
    print(internal)
    print(external)
    print(workbook)


if __name__ == "__main__":
    main()
